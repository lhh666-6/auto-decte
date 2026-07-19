from dataclasses import replace
from datetime import date
from hashlib import sha256
from pathlib import Path
from shutil import copyfile
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine

import app.adapters.export.xlsx as xlsx_adapter
from app.adapters.database.models import Base
from app.adapters.database.report_definition_repository_ds import (
    SqlAlchemyReportDefinitionRepository,
)
from app.adapters.database.repositories import SqlAlchemyFormRepository
from app.adapters.database.template_repository_ds import SqlAlchemyTemplateRepository
from app.adapters.export.xlsx import XlsxExporter
from app.adapters.storage.local import LocalEvidenceStorage
from app.application.export_forms import ExportForms
from app.application.import_forms import ImportForms
from app.application.query_forms import FormFilters, QueryForms, SearchResult
from app.application.review_forms import ReviewForms
from app.domain.models import (
    ExportStatus,
    Form,
    FormField,
    RecordStatus,
    RecordVersion,
    ReviewStatus,
)
from app.domain.templates_ds import (
    ExportTarget,
    FieldDefinition,
    FieldRules,
    PageSpec,
    Rect,
    TemplateVersion,
)
from app.modules.reporting.facade_ds import ReportingFacade
from app.modules.reporting.models_ds import (
    BUILTIN_REPORT_DEFINITIONS,
    AggregateOperation,
    ExportMapping,
    FixedCellMapping,
    FixedTableColumn,
    FixedTableMapping,
    ReportAggregate,
    ReportColumn,
    ReportDefinition,
    ReportDefinitionStatus,
    ReportKind,
)


def setup_confirmed_form(tmp_path: Path):  # type: ignore[no-untyped-def]
    engine = create_engine(f"sqlite:///{tmp_path / 'demo.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "evidence")
    )
    reviews = ReviewForms(repository, repository)
    image = tmp_path / "scan.png"
    image.write_bytes(b"image")
    evidence = imports.import_image(image, "FORM-0001", "T1", "1", "operator")
    reviews.confirm(
        "FORM-0001",
        0,
        {"employee_id": "E001", "total_quantity": 10, "qualified_quantity": 9},
        "reviewer",
        "confirmed",
        (evidence.file_id,),
    )
    service = ExportForms(repository, XlsxExporter(), QueryForms(repository))
    return service, reviews, repository


def test_xlsx_has_four_sheets_persisted_batch_and_reverse_trace(tmp_path: Path) -> None:
    service, _, repository = setup_confirmed_form(tmp_path)

    batch = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    workbook = load_workbook(batch.file_path, read_only=True)
    assert workbook.sheetnames == ["正式数据", "异常与复核", "汇总", "导出说明"]
    official_rows = list(workbook["正式数据"].iter_rows(values_only=True))
    headers, data = official_rows
    row = dict(zip(headers, data, strict=True))
    assert row["export_batch_id"] == batch.export_batch_id
    assert row["form_id"] == "FORM-0001"
    assert row["record_version"] == 1
    assert Path(batch.file_path).exists()
    assert len(batch.file_sha256) == 64
    assert repository.list_export_batches()[0].included_records == (("FORM-0001", 1),)
    assert repository.get_form("FORM-0001").export_status is ExportStatus.EXPORTED


def test_correction_after_export_requires_new_export_and_old_file_remains(tmp_path: Path) -> None:
    service, reviews, repository = setup_confirmed_form(tmp_path)
    first = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    reviews.confirm(
        "FORM-0001",
        1,
        {"employee_id": "E001", "total_quantity": 11, "qualified_quantity": 10},
        "reviewer",
        "correction",
        (),
    )

    assert repository.get_form("FORM-0001").export_status is ExportStatus.REEXPORT_REQUIRED
    assert Path(first.file_path).exists()
    second = service.export(
        "OUTPUT",
        FormFilters(),
        tmp_path / "exports",
        "finance",
        supersedes_batch_id=first.export_batch_id,
    )
    assert second.supersedes_batch_id == first.export_batch_id
    assert second.file_path != first.file_path
    assert Path(first.file_path).exists()
    assert Path(second.file_path).exists()


def test_preview_uses_stable_template_mapping_and_explains_unconfirmed_forms(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'preview.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "evidence")
    )
    reviews = ReviewForms(repository, repository)
    template = TemplateVersion.draft("TPL-T1-V1", "T1", 1, PageSpec.a4_portrait())
    template.add_field(
        FieldDefinition(
            "employee_id",
            "Display name must not become an export column",
            "text",
            "text_box",
            Rect(0.1, 0.1, 0.2, 0.05),
            template.page,
            export_target=ExportTarget("payroll.xlsx", "employees", "employee_code"),
        )
    )
    templates.add_version(template)
    for sequence in (1, 2):
        image = tmp_path / f"scan-{sequence}.png"
        image.write_bytes(f"image-{sequence}".encode())
        evidence = imports.import_image(image, f"FORM-{sequence:04d}", "T1", "1", "operator")
        if sequence == 1:
            reviews.confirm(
                "FORM-0001",
                0,
                {"employee_id": "E001"},
                "reviewer",
                "confirmed",
                (evidence.file_id,),
            )
    service = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        template_repository=templates,
    )

    preview = service.preview(FormFilters(), actor_id="finance")

    assert [(item.form_id, item.record_version) for item in preview.included] == [("FORM-0001", 1)]
    assert [(item.form_id, item.reason) for item in preview.excluded] == [
        ("FORM-0002", "NOT_CONFIRMED")
    ]
    assert [
        (
            item.template_id,
            item.template_version,
            item.field_key,
            item.workbook,
            item.worksheet,
            item.business_column,
        )
        for item in preview.mapping_snapshot
    ] == [("T1", "1", "employee_id", "payroll.xlsx", "employees", "employee_code")]
    assert repository.list_export_batches() == []
    assert repository.get_form("FORM-0001").export_status is ExportStatus.NOT_EXPORTED  # type: ignore[union-attr]


def test_preview_excludes_confirmed_record_that_fails_basic_final_validation(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'invalid-preview.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    template = TemplateVersion.draft("TPL-T1-V1", "T1", 1, PageSpec.a4_portrait())
    template.add_field(
        FieldDefinition(
            "employee_id",
            "Employee",
            "text",
            "text_box",
            Rect(0.1, 0.1, 0.2, 0.05),
            template.page,
            rules=FieldRules(required=True),
            export_target=ExportTarget("payroll.xlsx", "employees", "employee_code"),
        )
    )
    templates.add_version(template)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "invalid-evidence")
    )
    image = tmp_path / "invalid.png"
    image.write_bytes(b"invalid")
    evidence = imports.import_image(image, "FORM-INVALID", "T1", "1", "operator")
    ReviewForms(repository, repository).confirm(
        "FORM-INVALID", 0, {}, "reviewer", "confirmed", (evidence.file_id,)
    )
    service = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        template_repository=templates,
    )

    preview = service.preview(FormFilters())

    assert preview.included == ()
    assert [(item.form_id, item.reason) for item in preview.excluded] == [
        ("FORM-INVALID", "FINAL_VALIDATION_FAILED")
    ]
    assert preview.mapping_snapshot == ()

    with pytest.raises(ValueError, match="No exportable records after final validation"):
        service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    assert repository.list_export_batches() == []
    assert list((tmp_path / "exports").glob("*")) == []
    assert repository.get_form("FORM-INVALID").export_status is ExportStatus.NOT_EXPORTED  # type: ignore[union-attr]


def test_reporting_facade_preview_requires_template_repository(tmp_path: Path) -> None:
    _, _, repository = setup_confirmed_form(tmp_path)
    facade = ReportingFacade(repository, QueryForms(repository))

    with pytest.raises(RuntimeError, match="template repository is required for export preview"):
        facade.preview(FormFilters())


def test_reporting_facade_preview_explains_missing_template(tmp_path: Path) -> None:
    _, _, repository = setup_confirmed_form(tmp_path)
    templates = SqlAlchemyTemplateRepository(create_engine(f"sqlite:///{tmp_path / 'demo.db'}"))
    facade = ReportingFacade(repository, QueryForms(repository), template_repository=templates)

    preview = facade.preview(FormFilters())

    assert preview.included == ()
    assert [(item.form_id, item.reason) for item in preview.excluded] == [
        ("FORM-0001", "TEMPLATE_NOT_FOUND")
    ]


def test_preview_excludes_template_without_any_export_mapping(tmp_path: Path) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'unmapped-preview.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    templates.add_version(TemplateVersion.draft("TPL-EMPTY-V1", "EMPTY", 1, PageSpec.a4_portrait()))
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "unmapped-evidence")
    )
    image = tmp_path / "unmapped.png"
    image.write_bytes(b"unmapped")
    evidence = imports.import_image(image, "FORM-UNMAPPED", "EMPTY", "1", "operator")
    ReviewForms(repository, repository).confirm(
        "FORM-UNMAPPED", 0, {}, "reviewer", "confirmed", (evidence.file_id,)
    )
    service = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        template_repository=templates,
    )

    preview = service.preview(FormFilters())

    assert preview.included == ()
    assert [(item.form_id, item.reason) for item in preview.excluded] == [
        ("FORM-UNMAPPED", "NO_VALID_MAPPING")
    ]


def test_template_export_groups_mappings_and_uses_only_each_forms_template(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'mapped-export.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "mapped-evidence")
    )
    reviews = ReviewForms(repository, repository)

    template_one = TemplateVersion.draft("TPL-T1-V1", "T1", 1, PageSpec.a4_portrait())
    template_one.add_field(
        FieldDefinition(
            "employee_id",
            "This display name must never become a column",
            "text",
            "text_box",
            Rect(0.1, 0.1, 0.2, 0.05),
            template_one.page,
            export_target=ExportTarget("企业工资记录.xlsx", "人员", "员工编号"),
        )
    )
    template_one.add_field(
        FieldDefinition(
            "total_quantity",
            "Another ignored display name",
            "integer",
            "number",
            Rect(0.1, 0.2, 0.2, 0.05),
            template_one.page,
            export_target=ExportTarget("企业工资记录.xlsx", "产量", "总产量"),
        )
    )
    template_two = TemplateVersion.draft("TPL-T2-V1", "T2", 1, PageSpec.a4_portrait())
    template_two.add_field(
        FieldDefinition(
            "employee_id",
            "Also ignored",
            "text",
            "text_box",
            Rect(0.1, 0.1, 0.2, 0.05),
            template_two.page,
            export_target=ExportTarget("企业工资记录.xlsx", "人员", "外部人员编号"),
        )
    )
    templates.add_version(template_one)
    templates.add_version(template_two)

    for form_id, template_id, values in (
        ("FORM-T1", "T1", {"employee_id": "E001", "total_quantity": 10}),
        ("FORM-T2", "T2", {"employee_id": "EXT-2"}),
    ):
        image = tmp_path / f"{form_id}.png"
        image.write_bytes(form_id.encode())
        evidence = imports.import_image(image, form_id, template_id, "1", "operator")
        runtime_values: dict[str, object] = {}
        for field_name, value in values.items():
            field_id = f"{form_id}:RUNTIME:{field_name}"
            repository.add_form_field(
                FormField(
                    field_id=field_id,
                    form_id=form_id,
                    field_name=field_name,
                    source_region={"x": 0, "y": 0, "width": 1, "height": 1},
                )
            )
            runtime_values[field_id] = value
        reviews.confirm(
            form_id,
            0,
            runtime_values,
            "reviewer",
            "confirmed",
            (evidence.file_id,),
        )

    service = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        template_repository=templates,
    )

    batch = service.export("OUTPUT", FormFilters(), tmp_path / "exports", "finance")

    workbook = load_workbook(batch.file_path, read_only=True)
    assert workbook.sheetnames == ["人员", "产量"]
    people_rows = list(workbook["人员"].iter_rows(values_only=True))
    people = [dict(zip(people_rows[0], row, strict=True)) for row in people_rows[1:]]
    assert people == [
        {
            "export_batch_id": batch.export_batch_id,
            "form_id": "FORM-T1",
            "record_version": 1,
            "员工编号": "E001",
            "外部人员编号": None,
        },
        {
            "export_batch_id": batch.export_batch_id,
            "form_id": "FORM-T2",
            "record_version": 1,
            "员工编号": None,
            "外部人员编号": "EXT-2",
        },
    ]
    production_rows = list(workbook["产量"].iter_rows(values_only=True))
    assert production_rows == [
        ("export_batch_id", "form_id", "record_version", "总产量"),
        (batch.export_batch_id, "FORM-T1", 1, 10),
    ]
    assert "This display name must never become a column" not in people_rows[0]


def test_template_export_rejects_multiple_workbooks(tmp_path: Path) -> None:
    result = _search_result({"first": "one", "second": "two"})
    mappings = (
        ExportMapping("T1", "1", "first", "first.xlsx", "data", "first"),
        ExportMapping("T1", "1", "second", "second.xlsx", "data", "second"),
    )

    with pytest.raises(ValueError, match="one workbook"):
        XlsxExporter().write(
            tmp_path / "multiple.xlsx",
            "BATCH-1",
            "OUTPUT",
            [result],
            {},
            mappings=mappings,
        )


def test_template_export_rejects_case_insensitive_worksheet_collisions(
    tmp_path: Path,
) -> None:
    mappings = (
        ExportMapping("T1", "1", "first", "one.xlsx", "data", "first"),
        ExportMapping("T1", "1", "second", "one.xlsx", "DATA", "second"),
    )

    with pytest.raises(ValueError, match="worksheet names must be unique ignoring case"):
        XlsxExporter().write(
            tmp_path / "worksheet-collision.xlsx",
            "BATCH-1",
            "OUTPUT",
            [_search_result({"first": "one", "second": "two"})],
            {},
            mappings=mappings,
        )


def test_template_export_rejects_duplicate_business_column_for_template_sheet(
    tmp_path: Path,
) -> None:
    mappings = (
        ExportMapping("T1", "1", "first", "one.xlsx", "data", "duplicate"),
        ExportMapping("T1", "1", "second", "one.xlsx", "data", "duplicate"),
    )

    with pytest.raises(ValueError, match="business columns must be unique"):
        XlsxExporter().write(
            tmp_path / "column-collision.xlsx",
            "BATCH-1",
            "OUTPUT",
            [_search_result({"first": "one", "second": "two"})],
            {},
            mappings=mappings,
        )


@pytest.mark.parametrize(
    ("unsafe", "expected"),
    [
        ("=1+1", "'=1+1"),
        ("+SUM(A1:A2)", "'+SUM(A1:A2)"),
        ("-2+3", "'-2+3"),
        ("@cmd", "'@cmd"),
        ("\ufeff=hidden", "'=hidden"),
        ("#N/A", "#N/A"),
        ("#REF!", "#REF!"),
    ],
)
def test_template_export_neutralizes_formula_like_text(
    tmp_path: Path, unsafe: str, expected: str
) -> None:
    destination = tmp_path / "safe.xlsx"
    XlsxExporter().write(
        destination,
        "BATCH-SAFE",
        "OUTPUT",
        [_search_result({"value": unsafe})],
        {},
        mappings=(ExportMapping("T1", "1", "value", "safe.xlsx", "data", "value"),),
    )

    cell = load_workbook(destination, data_only=False)["data"]["D2"]
    assert cell.value == expected
    assert cell.data_type == "s"


@pytest.mark.parametrize(
    ("value", "expected_data_type"),
    [(42, "n"), (True, "b"), (date(2026, 7, 16), "d"), (None, "n")],
)
def test_template_export_preserves_non_text_cell_types(
    tmp_path: Path, value: object, expected_data_type: str
) -> None:
    destination = tmp_path / "typed.xlsx"
    XlsxExporter().write(
        destination,
        "BATCH-TYPED",
        "OUTPUT",
        [_search_result({"value": value})],
        {},
        mappings=(ExportMapping("T1", "1", "value", "typed.xlsx", "data", "value"),),
    )

    cell = load_workbook(destination, data_only=False)["data"]["D2"]
    assert cell.data_type == expected_data_type


def test_detail_report_definition_selects_and_sorts_controlled_columns(
    tmp_path: Path,
) -> None:
    definition = ReportDefinition(
        "DETAIL:1",
        "DETAIL",
        1,
        "工资明细",
        ReportKind.DETAIL,
        ReportDefinitionStatus.PUBLISHED,
        columns=(
            ReportColumn("employee_id", "员工编号"),
            ReportColumn("amount", "金额"),
        ),
        sort_by=("employee_id",),
        worksheet="工资明细",
    )
    first = _search_result({"employee_id": "E002", "amount": 20, "ignored": "x"})
    second = _search_result({"employee_id": "E001", "amount": 10, "ignored": "y"})
    destination = tmp_path / "detail.xlsx"

    XlsxExporter().write(
        destination,
        "BATCH-DETAIL",
        "DETAIL",
        [first, second],
        {},
        report_definition=definition,
    )

    rows = list(load_workbook(destination, read_only=True)["工资明细"].iter_rows(values_only=True))
    assert rows[0] == (
        "export_batch_id",
        "form_id",
        "record_version",
        "员工编号",
        "金额",
    )
    assert [row[3:] for row in rows[1:]] == [("E001", 10), ("E002", 20)]
    assert "ignored" not in rows[0]


def test_summary_report_definition_groups_and_aggregates_without_formulas(
    tmp_path: Path,
) -> None:
    definition = ReportDefinition(
        "SUMMARY:1",
        "SUMMARY",
        1,
        "员工汇总",
        ReportKind.SUMMARY,
        ReportDefinitionStatus.PUBLISHED,
        columns=(ReportColumn("employee_id", "员工编号"),),
        group_by=("employee_id",),
        aggregates=(
            ReportAggregate("amount", AggregateOperation.SUM, "工资合计"),
            ReportAggregate("amount", AggregateOperation.COUNT, "记录数"),
            ReportAggregate("quantity", AggregateOperation.AVERAGE, "平均数量"),
        ),
        sort_by=("employee_id",),
        worksheet="员工汇总",
    )
    destination = tmp_path / "summary.xlsx"

    XlsxExporter().write(
        destination,
        "BATCH-SUMMARY",
        "SUMMARY",
        [
            _search_result(
                {
                    "worker_number": "E002",
                    "calculated_wage": 5,
                    "qualified_quantity_1": 2,
                }
            ),
            _search_result(
                {
                    "worker_number": "E001",
                    "calculated_wage": 10,
                    "qualified_quantity_1": 4,
                }
            ),
            _search_result(
                {
                    "worker_number": "E001",
                    "calculated_wage": 15,
                    "qualified_quantity_1": 6,
                }
            ),
        ],
        {},
        report_definition=definition,
    )

    rows = list(load_workbook(destination, read_only=True)["员工汇总"].iter_rows(values_only=True))
    assert rows == [
        ("export_batch_id", "员工编号", "工资合计", "记录数", "平均数量"),
        ("BATCH-SUMMARY", "E001", 25, 2, 5),
        ("BATCH-SUMMARY", "E002", 5, 1, 2),
    ]


def test_two_builtin_summary_definitions_generate_distinct_business_outputs(
    tmp_path: Path,
) -> None:
    definitions = {item.report_key: item for item in BUILTIN_REPORT_DEFINITIONS}
    rows = [
        _search_result(
            {
                "worker_number": "E001",
                "worker_name": "张三",
                "work_order_number": "1001",
                "calculated_wage": 10,
                "qualified_quantity_1": 4,
            }
        ),
        _search_result(
            {
                "worker_number": "E001",
                "worker_name": "张三",
                "work_order_number": "1001",
                "calculated_wage": 15,
                "qualified_quantity_1": 6,
            }
        ),
    ]
    employee_path = tmp_path / "employee-summary.xlsx"
    work_order_path = tmp_path / "work-order-summary.xlsx"

    XlsxExporter().write(
        employee_path,
        "BATCH-EMPLOYEE",
        "EMPLOYEE_PAYROLL_SUMMARY",
        rows,
        {},
        report_definition=definitions["EMPLOYEE_PAYROLL_SUMMARY"],
    )
    XlsxExporter().write(
        work_order_path,
        "BATCH-WORK-ORDER",
        "WORK_ORDER_OUTPUT_SUMMARY",
        rows,
        {},
        report_definition=definitions["WORK_ORDER_OUTPUT_SUMMARY"],
    )

    employee_rows = list(
        load_workbook(employee_path, read_only=True)["员工工资汇总"].iter_rows(values_only=True)
    )
    work_order_rows = list(
        load_workbook(work_order_path, read_only=True)["工单产量汇总"].iter_rows(values_only=True)
    )
    assert employee_rows[1][1:] == ("E001", "张三", 25)
    assert work_order_rows[1][1:] == ("1001", 10)


def test_fixed_report_copies_trusted_template_and_preserves_layout(tmp_path: Path) -> None:
    assert hasattr(xlsx_adapter, "FixedTemplateAsset")
    asset_path = (
        Path(xlsx_adapter.__file__).parents[2]
        / "modules"
        / "reporting"
        / "assets"
        / "legacy_timekeeping_daily.xlsx"
    )
    asset_hash = sha256(asset_path.read_bytes()).hexdigest()
    definition = ReportDefinition(
        "TIMEKEEPING_DAILY_FIXED:1",
        "TIMEKEEPING_DAILY_FIXED",
        1,
        "计时工日工资表（原格式）",
        ReportKind.FIXED,
        ReportDefinitionStatus.PUBLISHED,
        worksheet="计时",
        fixed_template_key="LEGACY_TIMEKEEPING_DAILY",
        fixed_template_sha256=asset_hash,
        fixed_cells=(FixedCellMapping("H2", "work_date"),),
        fixed_table=FixedTableMapping(
            5,
            8,
            (
                FixedTableColumn("A", "employee_name"),
                FixedTableColumn("D", "hours"),
                FixedTableColumn("E", "assessment"),
            ),
        ),
    )
    destination = tmp_path / "fixed.xlsx"

    XlsxExporter().write(
        destination,
        "BATCH-FIXED",
        definition.report_key,
        [
            _search_result(
                {
                    "work_date": "2026-07-20",
                    "worker_name": "张三",
                    "effective_hours_1": 4,
                    "effective_hours_2": 3.5,
                    "assessment": "=A",
                }
            )
        ],
        {},
        report_definition=definition,
    )

    assert sha256(asset_path.read_bytes()).hexdigest() == asset_hash
    source_worksheet = load_workbook(asset_path, read_only=False)["计时"]
    workbook = load_workbook(destination, read_only=False)
    worksheet = workbook["计时"]
    assert worksheet["A1"].value == "计时工日工资考核表"
    assert "A1:J1" in {str(item) for item in worksheet.merged_cells.ranges}
    assert worksheet["H2"].value == "2026-07-20"
    assert worksheet["A5"].value == "张三"
    assert worksheet["D5"].value == 7.5
    assert worksheet["E5"].value == "'=A"
    assert worksheet.page_setup.orientation == source_worksheet.page_setup.orientation
    assert worksheet.page_setup.paperSize == source_worksheet.page_setup.paperSize
    assert worksheet.column_dimensions["A"].width == source_worksheet.column_dimensions["A"].width
    assert worksheet.row_dimensions[1].height == source_worksheet.row_dimensions[1].height


def test_fixed_report_rejects_formula_template_hash_mismatch_and_row_overflow(
    tmp_path: Path,
) -> None:
    assert hasattr(xlsx_adapter, "FixedTemplateAsset")
    template = tmp_path / "unsafe.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "计时"
    worksheet["A1"] = "=1+1"
    workbook.save(template)
    digest = sha256(template.read_bytes()).hexdigest()
    definition = ReportDefinition(
        "UNSAFE_FIXED:1",
        "UNSAFE_FIXED",
        1,
        "不安全固定表",
        ReportKind.FIXED,
        ReportDefinitionStatus.PUBLISHED,
        worksheet="计时",
        fixed_template_key="UNSAFE_TEMPLATE",
        fixed_template_sha256=digest,
        fixed_cells=(FixedCellMapping("B1", "work_date"),),
        fixed_table=FixedTableMapping(
            2,
            1,
            (FixedTableColumn("A", "employee_name"),),
        ),
    )
    exporter = XlsxExporter(
        fixed_template_assets={
            "UNSAFE_TEMPLATE": xlsx_adapter.FixedTemplateAsset("UNSAFE_TEMPLATE", template, digest)
        }
    )

    with pytest.raises(ValueError, match="formulas"):
        exporter.write(
            tmp_path / "formula.xlsx",
            "BATCH",
            definition.report_key,
            [_search_result({"worker_name": "张三"})],
            {},
            report_definition=definition,
        )
    with pytest.raises(ValueError, match="hash"):
        XlsxExporter(
            fixed_template_assets={
                "UNSAFE_TEMPLATE": xlsx_adapter.FixedTemplateAsset(
                    "UNSAFE_TEMPLATE", template, "0" * 64
                )
            }
        ).write(
            tmp_path / "hash.xlsx",
            "BATCH",
            definition.report_key,
            [_search_result({"worker_name": "张三"})],
            {},
            report_definition=definition,
        )
    safe_template = tmp_path / "safe.xlsx"
    workbook["计时"]["A1"] = "标题"
    workbook.save(safe_template)
    safe_digest = sha256(safe_template.read_bytes()).hexdigest()
    safe_definition = replace(definition, fixed_template_sha256=safe_digest)
    safe_exporter = XlsxExporter(
        fixed_template_assets={
            "UNSAFE_TEMPLATE": xlsx_adapter.FixedTemplateAsset(
                "UNSAFE_TEMPLATE", safe_template, safe_digest
            )
        }
    )
    with pytest.raises(ValueError, match="at most 1 rows"):
        safe_exporter.write(
            tmp_path / "overflow.xlsx",
            "BATCH",
            safe_definition.report_key,
            [
                _search_result({"worker_name": "张三"}),
                _search_result({"worker_name": "李四"}),
            ],
            {},
            report_definition=safe_definition,
        )


@pytest.mark.parametrize(
    ("member", "message"),
    [
        ("xl/vbaProject.bin", "macros"),
        ("xl/externalLinks/externalLink1.xml", "external links"),
    ],
)
def test_fixed_report_rejects_macro_and_external_link_packages(
    tmp_path: Path,
    member: str,
    message: str,
) -> None:
    base = tmp_path / "base.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "计时"
    workbook.save(base)
    unsafe = tmp_path / f"unsafe-{message.replace(' ', '-')}.xlsx"
    copyfile(base, unsafe)
    with ZipFile(unsafe, "a") as package:
        package.writestr(member, b"unsafe")
    digest = sha256(unsafe.read_bytes()).hexdigest()
    definition = ReportDefinition(
        "PACKAGE_FIXED:1",
        "PACKAGE_FIXED",
        1,
        "包安全测试",
        ReportKind.FIXED,
        ReportDefinitionStatus.PUBLISHED,
        worksheet="计时",
        fixed_template_key="PACKAGE_TEMPLATE",
        fixed_template_sha256=digest,
        fixed_cells=(FixedCellMapping("A1", "work_date"),),
    )
    exporter = XlsxExporter(
        fixed_template_assets={
            "PACKAGE_TEMPLATE": xlsx_adapter.FixedTemplateAsset("PACKAGE_TEMPLATE", unsafe, digest)
        }
    )

    with pytest.raises(ValueError, match=message):
        exporter.write(
            tmp_path / "output.xlsx",
            "BATCH",
            definition.report_key,
            [_search_result({"work_date": "2026-07-20"})],
            {},
            report_definition=definition,
        )


def test_export_resolves_published_definition_and_snapshots_exact_version(
    tmp_path: Path,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'report-export.db'}")
    Base.metadata.create_all(engine)
    repository = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    reports = SqlAlchemyReportDefinitionRepository(engine)
    template = TemplateVersion.draft("TPL-T1-V1", "T1", 1, PageSpec.a4_portrait())
    for index, field_name in enumerate(("employee_id", "amount")):
        template.add_field(
            FieldDefinition(
                field_name,
                field_name,
                "text" if field_name == "employee_id" else "decimal",
                "text_box" if field_name == "employee_id" else "number_box",
                Rect(0.1, 0.1 + index * 0.1, 0.2, 0.05),
                template.page,
                export_target=ExportTarget("legacy.xlsx", "data", field_name),
            )
        )
    templates.add_version(template)
    definition = ReportDefinition(
        "EMPLOYEE_TOTAL:2",
        "EMPLOYEE_TOTAL",
        2,
        "员工合计",
        ReportKind.SUMMARY,
        ReportDefinitionStatus.PUBLISHED,
        columns=(ReportColumn("employee_id", "员工编号"),),
        group_by=("employee_id",),
        aggregates=(ReportAggregate("amount", AggregateOperation.SUM, "金额合计"),),
        worksheet="员工合计",
    )
    reports.add(definition)
    imports = ImportForms(
        repository, repository, repository, LocalEvidenceStorage(tmp_path / "report-evidence")
    )
    reviews = ReviewForms(repository, repository)
    for form_id, employee_id, amount in (
        ("FORM-1", "E001", 10),
        ("FORM-2", "E001", 15),
    ):
        image = tmp_path / f"{form_id}.png"
        image.write_bytes(form_id.encode())
        evidence = imports.import_image(image, form_id, "T1", "1", "operator")
        reviews.confirm(
            form_id,
            0,
            {"employee_id": employee_id, "amount": amount},
            "reviewer",
            "confirmed",
            (evidence.file_id,),
        )
    service = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        template_repository=templates,
        report_definition_repository=reports,
    )

    batch = service.export(
        "EMPLOYEE_TOTAL",
        FormFilters(),
        tmp_path / "exports",
        "finance",
        report_definition_id="EMPLOYEE_TOTAL:2",
    )

    worksheet = load_workbook(batch.file_path, read_only=True)["员工合计"]
    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[1][1:] == ("E001", 25)
    snapshot = batch.template_snapshot["report_definition"]
    assert snapshot["definition_id"] == "EMPLOYEE_TOTAL:2"
    assert snapshot["version"] == 2
    assert snapshot["aggregates"][0]["operation"] == "SUM"


def test_export_rejects_unknown_or_unpublished_report_definition(tmp_path: Path) -> None:
    _, _, repository = setup_confirmed_form(tmp_path)
    reports = SqlAlchemyReportDefinitionRepository(
        create_engine(f"sqlite:///{tmp_path / 'demo.db'}")
    )
    reports.add(
        ReportDefinition(
            "DRAFT_REPORT:1",
            "DRAFT_REPORT",
            1,
            "草稿报表",
            ReportKind.DETAIL,
            ReportDefinitionStatus.DRAFT,
            columns=(ReportColumn("employee_id", "员工编号"),),
        )
    )
    controlled = ExportForms(
        repository,
        XlsxExporter(),
        QueryForms(repository),
        report_definition_repository=reports,
    )

    with pytest.raises(KeyError, match="Unknown report definition"):
        controlled.export(
            "UNKNOWN",
            FormFilters(),
            tmp_path / "unknown",
            "finance",
            report_definition_id="UNKNOWN:1",
        )
    with pytest.raises(ValueError, match="published"):
        controlled.export(
            "DRAFT_REPORT",
            FormFilters(),
            tmp_path / "draft",
            "finance",
            report_definition_id="DRAFT_REPORT:1",
        )


def _search_result(values: dict[str, object]) -> SearchResult:
    form = Form(
        "FORM-SAFE",
        "T1",
        "1",
        review_status=ReviewStatus.CONFIRMED,
        current_record_version=1,
    )
    record = RecordVersion(
        "RECORD-SAFE",
        form.form_id,
        1,
        RecordStatus.CONFIRMED,
        values,
        confirmed_by="reviewer",
    )
    return SearchResult(form, record)
