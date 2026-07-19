"""Stable Excel layouts for the reviewed payroll templates."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine

from app.adapters.database.models import Base
from app.adapters.database.repositories import SqlAlchemyFormRepository
from app.adapters.database.template_repository_ds import SqlAlchemyTemplateRepository
from app.adapters.export.xlsx import XlsxExporter
from app.adapters.storage.local import LocalEvidenceStorage
from app.application.export_forms import ExportForms
from app.application.import_forms import ImportForms
from app.application.query_forms import FormFilters, QueryForms, SearchResult
from app.application.review_forms import ReviewForms
from app.domain.models import Form, RecordStatus, RecordVersion, ReviewStatus
from app.domain.templates_ds import TemplateVersion
from app.modules.reporting.models_ds import ExportMapping
from app.modules.templates.payroll_profiles_ds import (
    reviewed_payroll_export_seed_templates,
)
from app.modules.templates.seed_templates_ds import legacy_payroll_seed_templates

MAIN_COLUMNS = [
    "日期",
    "班次",
    "工号",
    "姓名",
    "班组",
    "工单号",
    "考评合格",
    "需要改进",
    "考评不合格",
    "事实说明",
    "员工签字",
    "质量签字",
    "主管签字",
]
TRACE_COLUMNS = ["export_batch_id", "form_id", "record_version"]


@pytest.mark.parametrize(
    "template_key",
    [
        "PAYROLL_TIMEKEEPING_DAILY",
        "PAYROLL_STEAMING_DAILY",
        "PAYROLL_SHEET_CUTTING_DAILY",
    ],
)
def test_three_payroll_masters_export_stable_main_detail_columns_and_excel_types(
    tmp_path: Path,
    template_key: str,
) -> None:
    template = _reviewed_template(template_key)
    batch = _export_confirmed_template(tmp_path, template)
    workbook = load_workbook(batch.file_path, data_only=False)

    assert workbook.sheetnames == ["工资主记录", "业务明细"]
    main = workbook["工资主记录"]
    detail = workbook["业务明细"]
    assert [cell.value for cell in main[1]] == TRACE_COLUMNS + MAIN_COLUMNS
    detail_fields = [
        field
        for field in template.fields
        if field.export_target is not None and field.export_target.worksheet == "业务明细"
    ]
    assert [cell.value for cell in detail[1]] == TRACE_COLUMNS + [
        field.display_name for field in detail_fields
    ]

    main_row = _row_by_header(main)
    assert main_row["工号"] == "W-001"
    assert main_row["姓名"] == "测试员工"
    assert main_row["考评合格"] is True
    assert main_row["需要改进"] is False
    assert main_row["事实说明"] == "无异常"
    assert main_row["员工签字"] == "员工签字"
    assert main_row["质量签字"] == "质量签字"
    assert main_row["主管签字"] == "主管签字"

    for field in template.fields:
        target = field.export_target
        assert target is not None
        worksheet = workbook[target.worksheet]
        cell = _cell_by_header(worksheet, target.business_column)
        expected_type = {
            "integer": "n",
            "decimal": "n",
            "boolean": "b",
        }.get(field.data_type, "s")
        assert cell.data_type == expected_type, (template_key, field.field_key, cell.value)

    assert [item["field_key"] for item in batch.mapping_snapshot] == [
        field.field_key for field in template.fields
    ]


def test_fixed_production_grid_expands_each_form_as_linked_main_and_detail_rows(
    tmp_path: Path,
) -> None:
    template = _reviewed_template("PAYROLL_SHEET_CUTTING_DAILY")
    mappings = _mappings(template)
    results = [
        _search_result(template, "FORM-CUT-1", 1),
        _search_result(template, "FORM-CUT-2", 2),
    ]
    destination = tmp_path / "fixed-grid.xlsx"

    XlsxExporter().write(
        destination,
        "BATCH-FIXED",
        "PAYROLL",
        results,
        {},
        mappings=mappings,
    )

    workbook = load_workbook(destination, read_only=True)
    main_rows = list(workbook["工资主记录"].iter_rows(values_only=True))
    detail_rows = list(workbook["业务明细"].iter_rows(values_only=True))
    assert len(main_rows) == len(detail_rows) == 3
    assert [row[:3] for row in main_rows[1:]] == [
        ("BATCH-FIXED", "FORM-CUT-1", 1),
        ("BATCH-FIXED", "FORM-CUT-2", 2),
    ]
    assert [row[:3] for row in detail_rows[1:]] == [
        ("BATCH-FIXED", "FORM-CUT-1", 1),
        ("BATCH-FIXED", "FORM-CUT-2", 2),
    ]


def test_legacy_v1_template_keeps_its_original_single_sheet_mapping(
    tmp_path: Path,
) -> None:
    template = legacy_payroll_seed_templates()[0]
    destination = tmp_path / "legacy-v1.xlsx"

    XlsxExporter().write(
        destination,
        "BATCH-LEGACY",
        "PAYROLL",
        [_search_result(template, "FORM-LEGACY", 1)],
        {},
        mappings=_mappings(template),
    )

    workbook = load_workbook(destination, read_only=True)
    assert workbook.sheetnames == ["计时工资"]
    rows = list(workbook["计时工资"].iter_rows(values_only=True))
    assert list(rows[0]) == TRACE_COLUMNS + [field.field_key for field in template.fields]
    assert rows[1][:3] == ("BATCH-LEGACY", "FORM-LEGACY", 1)


def _reviewed_template(template_key: str) -> TemplateVersion:
    return next(
        template
        for template in reviewed_payroll_export_seed_templates()
        if template.template_key == template_key
    )


def _export_confirmed_template(tmp_path: Path, template: TemplateVersion):  # type: ignore[no-untyped-def]
    engine = create_engine(f"sqlite:///{tmp_path / 'payroll-export.db'}")
    Base.metadata.create_all(engine)
    forms = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    templates.add_version(template)
    imports = ImportForms(
        forms,
        forms,
        forms,
        LocalEvidenceStorage(tmp_path / "evidence"),
    )
    image = tmp_path / "payroll.png"
    image.write_bytes(b"payroll")
    evidence = imports.import_image(
        image,
        "FORM-PAYROLL",
        template.template_key,
        str(template.version),
        "operator",
    )
    ReviewForms(forms, forms).confirm(
        "FORM-PAYROLL",
        0,
        _values(template, 1),
        "reviewer",
        "confirmed",
        (evidence.file_id,),
    )
    return ExportForms(
        forms,
        XlsxExporter(),
        QueryForms(forms),
        template_repository=templates,
    ).export("PAYROLL", FormFilters(), tmp_path / "exports", "finance")


def _search_result(
    template: TemplateVersion,
    form_id: str,
    record_version: int,
) -> SearchResult:
    form = Form(
        form_id,
        template.template_key,
        str(template.version),
        review_status=ReviewStatus.CONFIRMED,
        current_record_version=record_version,
    )
    record = RecordVersion(
        f"RECORD-{form_id}",
        form_id,
        record_version,
        RecordStatus.CONFIRMED,
        _values(template, record_version),
        confirmed_by="reviewer",
    )
    return SearchResult(form, record)


def _values(template: TemplateVersion, marker: int) -> dict[str, object]:
    fixed = {
        "work_date": "2026-07-19",
        "shift": "白班",
        "worker_number": "W-001",
        "worker_name": "测试员工",
        "team_name": "一班",
        "work_order_number": f"WO-{marker:03d}",
        "assessment_passed": True,
        "assessment_improvement": False,
        "assessment_failed": False,
        "facts_description": "无异常",
        "worker_signature": "员工签字",
        "quality_signature": "质量签字",
        "supervisor_signature": "主管签字",
    }
    values: dict[str, object] = {}
    for index, field in enumerate(template.fields, start=1):
        if field.field_key in fixed:
            values[field.field_key] = fixed[field.field_key]
        elif field.data_type == "integer":
            values[field.field_key] = marker * 10 + index
        elif field.data_type == "decimal":
            values[field.field_key] = marker * 10 + index + 0.5
        elif field.data_type == "boolean":
            values[field.field_key] = False
        else:
            values[field.field_key] = f"{field.display_name}-{marker}"
    return values


def _mappings(template: TemplateVersion) -> tuple[ExportMapping, ...]:
    return tuple(
        ExportMapping(
            template.template_key,
            str(template.version),
            field.field_key,
            field.export_target.workbook,
            field.export_target.worksheet,
            field.export_target.business_column,
        )
        for field in template.fields
        if field.export_target is not None
    )


def _row_by_header(worksheet) -> dict[str, object]:  # type: ignore[no-untyped-def]
    return {cell.value: worksheet.cell(row=2, column=cell.column).value for cell in worksheet[1]}


def _cell_by_header(worksheet, header: str):  # type: ignore[no-untyped-def]
    column = next(cell.column for cell in worksheet[1] if cell.value == header)
    return worksheet.cell(row=2, column=column)
