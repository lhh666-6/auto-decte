"""Digital paper-loop acceptance for the three representative payroll masters."""

from __future__ import annotations

import re
from pathlib import Path

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import create_engine

from app.adapters.database.models import Base
from app.adapters.database.repositories import SqlAlchemyFormRepository
from app.adapters.database.template_repository_ds import SqlAlchemyTemplateRepository
from app.adapters.export.xlsx import XlsxExporter
from app.adapters.recognition.opencv import OpenCvImagePipeline
from app.adapters.storage.local import LocalEvidenceStorage
from app.adapters.templates.print_renderer_ds import TemplatePrintRenderer
from app.application.export_forms import ExportForms
from app.application.import_forms import ImportForms
from app.application.query_forms import FormFilters, QueryForms
from app.application.recognize_forms import RecognizeForms
from app.application.review_forms import ReviewForms
from app.domain.templates_ds import PaperEntryMode, TemplateArtifact, TemplateVersion
from app.modules.templates.payroll_profiles_ds import reviewed_payroll_export_seed_templates

REPRESENTATIVE_KEYS = (
    "PAYROLL_TIMEKEEPING_DAILY",
    "PAYROLL_STEAMING_DAILY",
    "PAYROLL_SHEET_CUTTING_DAILY",
)
TRACE_EVENTS = {"IMPORT", "CLASSIFY", "NORMALIZE", "CROP_FIELDS", "CONFIRM", "EXPORT"}


def _artifact(artifacts: tuple[TemplateArtifact, ...], kind: str) -> TemplateArtifact:
    return next(item for item in artifacts if item.kind == kind)


def _templates() -> tuple[TemplateVersion, ...]:
    by_key = {
        template.template_key: template for template in reviewed_payroll_export_seed_templates()
    }
    return tuple(by_key[key] for key in REPRESENTATIVE_KEYS)


def _values(template: TemplateVersion, marker: int) -> dict[str, object]:
    fixed: dict[str, object] = {
        "work_date": "2026-07-19",
        "shift": "白班",
        "worker_number": f"W-{marker:03d}",
        "worker_name": "验收员工",
        "team_name": "一班",
        "work_order_number": f"WO-{marker:03d}",
        "assessment_passed": True,
        "assessment_improvement": False,
        "assessment_failed": False,
        "facts_description": "模拟纸面闭环验收",
        "worker_signature": "员工签字",
        "quality_signature": "质量签字",
        "supervisor_signature": "主管签字",
    }
    values: dict[str, object] = {}
    for index, field in enumerate(template.fields, start=1):
        if field.field_key in fixed:
            values[field.field_key] = fixed[field.field_key]
        elif field.rules.allowed_values:
            values[field.field_key] = field.rules.allowed_values[0]
        elif field.data_type == "integer":
            values[field.field_key] = max(int(field.rules.minimum_value or 0), marker * 10 + index)
        elif field.data_type == "decimal":
            values[field.field_key] = max(float(field.rules.minimum_value or 0), marker + 0.5)
        elif field.data_type == "boolean":
            values[field.field_key] = False
        else:
            values[field.field_key] = f"{field.display_name}-{marker}"
    return values


def _simulate_filled_photo(
    image: np.ndarray, template: TemplateVersion
) -> tuple[np.ndarray, str, tuple[float, float]]:
    """Add a visible field mark, mild perspective/rotation and an indoor-light shadow."""
    target = next(
        field
        for field in template.fields
        if field.paper_entry_mode is PaperEntryMode.DIGIT_BOXES
    )
    height, width = image.shape[:2]
    expected = (
        (target.region.x + target.region.width * 0.5) * width,
        (target.region.y + target.region.height * 0.76) * height,
    )
    radius = max(8, round(min(target.region.width * width, target.region.height * height) * 0.08))
    filled = image.copy()
    cv2.circle(filled, (round(expected[0]), round(expected[1])), radius, (255, 0, 255), -1)

    pad = max(90, min(width, height) // 18)
    output_width = width + pad * 2
    output_height = height + pad * 2
    source = np.float32([[0, 0], [width - 1, 0], [width - 1, height - 1], [0, height - 1]])
    destination = np.float32(
        [
            [pad + 24, pad + 10],
            [pad + width - 22, pad - 18],
            [pad + width + 18, pad + height + 22],
            [pad - 16, pad + height - 12],
        ]
    )
    transform = cv2.getPerspectiveTransform(source, destination)
    photographed = cv2.warpPerspective(
        filled,
        transform,
        (output_width, output_height),
        borderValue=(235, 235, 235),
    )
    shadow = np.linspace(0.86, 1.0, output_width, dtype=np.float32)[None, :, None]
    photographed = np.clip(photographed.astype(np.float32) * shadow, 0, 255).astype(np.uint8)
    return photographed, target.field_key, expected


def _assert_print_size(image_path: Path, template: TemplateVersion) -> None:
    with Image.open(image_path) as image:
        dpi_x, dpi_y = image.info["dpi"]
        width_mm = image.width / dpi_x * 25.4
        height_mm = image.height / dpi_y * 25.4
    assert abs(width_mm - template.page.width_mm) <= 1.0
    assert abs(height_mm - template.page.height_mm) <= 1.0


def _assert_pdf_size(pdf_path: Path, template: TemplateVersion) -> None:
    match = re.search(rb"/MediaBox \[ 0 0 ([0-9.]+) ([0-9.]+) \]", pdf_path.read_bytes())
    assert match is not None
    width_mm = float(match.group(1)) / 72 * 25.4
    height_mm = float(match.group(2)) / 72 * 25.4
    assert abs(width_mm - template.page.width_mm) <= 1.0
    assert abs(height_mm - template.page.height_mm) <= 1.0


def _assert_written_mark_alignment(
    corrected: np.ndarray,
    template: TemplateVersion,
    field_key: str,
    expected: tuple[float, float],
) -> None:
    blue, green, red = cv2.split(corrected)
    mask = (blue > 150) & (red > 150) & (green < 100)
    ys, xs = np.nonzero(mask)
    assert len(xs) > 20
    actual = (float(xs.mean()), float(ys.mean()))
    field = next(item for item in template.fields if item.field_key == field_key)
    short_edge = min(
        field.region.width * template.page.canonical_width_px,
        field.region.height * template.page.canonical_height_px,
    )
    deviation = float(np.hypot(actual[0] - expected[0], actual[1] - expected[1]))
    assert deviation <= short_edge * 0.10


def _assert_two_up_independent_identity(
    renderer: TemplatePrintRenderer, template: TemplateVersion
) -> None:
    preview = renderer.compose_imposition(
        template,
        print_batch="PB20260719A",
        first_sequence=41,
    )
    midpoint = preview.width // 2
    pipeline = OpenCvImagePipeline()
    left = cv2.cvtColor(
        np.asarray(preview.crop((0, 0, midpoint, preview.height))),
        cv2.COLOR_RGB2BGR,
    )
    right = cv2.cvtColor(
        np.asarray(preview.crop((midpoint, 0, preview.width, preview.height))),
        cv2.COLOR_RGB2BGR,
    )
    left_payloads = set(pipeline.read_qr_payloads(left))
    right_payloads = set(pipeline.read_qr_payloads(right))
    assert renderer.sheet_payload("PB20260719A", 41) in left_payloads
    assert renderer.sheet_payload("PB20260719A", 42) in right_payloads


@pytest.mark.parametrize("template", _templates(), ids=REPRESENTATIVE_KEYS)
def test_representative_template_completes_simulated_paper_to_excel_loop(
    tmp_path: Path,
    template: TemplateVersion,
) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'acceptance.db'}")
    Base.metadata.create_all(engine)
    forms = SqlAlchemyFormRepository(engine)
    templates = SqlAlchemyTemplateRepository(engine)
    templates.add_version(template)
    storage = LocalEvidenceStorage(tmp_path / "evidence")
    renderer = TemplatePrintRenderer(tmp_path / "print")
    pipeline = OpenCvImagePipeline()

    artifacts = renderer.render(template, print_batch="PB20260719A", sequence=10)
    png_path = Path(_artifact(artifacts, "PRINT_PNG").internal_uri)
    pdf_path = Path(_artifact(artifacts, "PRINT_PDF").internal_uri)
    assert pdf_path.read_bytes().startswith(b"%PDF")
    _assert_print_size(png_path, template)
    _assert_pdf_size(pdf_path, template)
    if template.template_key == "PAYROLL_STEAMING_DAILY":
        assert _artifact(artifacts, "PRINT_IMPOSED_PDF")
        _assert_two_up_independent_identity(renderer, template)

    rendered = cv2.imread(str(png_path))
    photographed, marked_field_key, expected_mark = _simulate_filled_photo(rendered, template)
    photo_path = tmp_path / f"{template.template_key}-photo.png"
    assert cv2.imwrite(str(photo_path), photographed)

    form_id = f"FORM-{template.template_key}"
    original = ImportForms(forms, forms, forms, storage).import_image(
        photo_path,
        form_id,
        "UNKNOWN",
        "0",
        "local-operator",
    )
    recognition = RecognizeForms(forms, forms, forms, storage, pipeline, templates)
    classification = recognition.classify_image(form_id, photographed)
    assert classification.template_reference is not None
    assert classification.source == "SHEET_QR"
    assert forms.get_form(form_id).template_version == str(template.version)  # type: ignore[union-attr]

    corrected_evidence, corrected = recognition.correct_and_record_template_canvas(
        form_id,
        photographed,
        width=template.page.canonical_width_px,
        height=template.page.canonical_height_px,
        canonical_dpi=template.page.canonical_dpi,
    )
    _assert_written_mark_alignment(corrected, template, marked_field_key, expected_mark)
    crops = recognition.record_template_field_crops(form_id, corrected, template)
    assert set(crops) == {field.field_key for field in template.fields}
    marked_crop = cv2.imread(str(tmp_path / "evidence" / crops[marked_field_key].uri))
    blue, green, red = cv2.split(marked_crop)
    assert int(np.count_nonzero((blue > 150) & (red > 150) & (green < 100))) > 20

    ReviewForms(forms, forms).confirm(
        form_id,
        0,
        _values(template, 1),
        "local-operator",
        "纸面模拟验收后人工复核",
        (original.file_id, corrected_evidence.file_id, crops[marked_field_key].file_id),
    )
    queries = QueryForms(forms)
    batch = ExportForms(
        forms,
        XlsxExporter(),
        queries,
        template_repository=templates,
    ).export(
        "PAYROLL",
        FormFilters(form_id=form_id),
        tmp_path / "exports",
        "local-operator",
    )

    workbook = load_workbook(batch.file_path, read_only=True, data_only=False)
    assert workbook.sheetnames == ["工资主记录", "业务明细"]
    assert list(workbook["工资主记录"].iter_rows(values_only=True))[1][1] == form_id
    trace = queries.trace(form_id)
    assert TRACE_EVENTS <= {event.event_type for event in trace.audits}
    assert trace.versions[-1].values["worker_name"] == "验收员工"
