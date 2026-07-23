"""Phase 6 external workbook, mapping, export and lineage acceptance tests."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine

from app.adapters.database.models import Base
from app.modules.report_templates.service_ds import ReportTemplateError, ReportTemplateService


@pytest.fixture
def reports(tmp_path):  # type: ignore[no-untyped-def]
    engine = create_engine(f"sqlite:///{(tmp_path / 'reports.db').as_posix()}")
    Base.metadata.create_all(engine)
    return ReportTemplateService(engine)


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工资明细"
    sheet["A1"] = "员工编号"
    sheet["B1"] = "工资"
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_xlsx_upload_records_safe_structure_and_hash(reports: ReportTemplateService) -> None:
    template = reports.upload_template(
        filename="工资模板.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=_xlsx(),
        actor_id="admin-1",
    )
    assert template["format"] == "XLSX"
    assert template["status"] == "ANALYZED"
    assert template["structure"]["sheets"][0]["name"] == "工资明细"
    assert len(str(template["file_hash"])) == 64
    assert len(str(template["structure_hash"])) == 64


def test_invalid_signature_and_formula_risk_are_rejected(
    reports: ReportTemplateService,
) -> None:
    with pytest.raises(ReportTemplateError, match="FILE_SIGNATURE_INVALID"):
        reports.upload_template(
            filename="伪造.xls",
            mime_type="application/vnd.ms-excel",
            content=b"not-an-ole-file",
            actor_id="admin-1",
        )

    workbook = Workbook()
    workbook.active["A1"] = '=DDE("cmd","/c calc",0)'
    target = BytesIO()
    workbook.save(target)
    with pytest.raises(ReportTemplateError, match="DANGEROUS_FORMULA"):
        reports.upload_template(
            filename="危险.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=target.getvalue(),
            actor_id="admin-1",
        )


def test_confirmed_mapping_exports_watermarked_data_and_cell_lineage(
    reports: ReportTemplateService,
) -> None:
    template = reports.upload_template(
        filename="工资模板.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=_xlsx(),
        actor_id="admin-1",
    )
    mapping = reports.create_mapping(
        template_version_id=str(template["template_version_id"]),
        mapping_json={
            "sheet": "工资明细",
            "start_row": 2,
            "columns": [
                {"column": 1, "source_field": "employee_code"},
                {"column": 2, "source_field": "amount"},
            ],
        },
        actor_id="finance-1",
    )
    with pytest.raises(ReportTemplateError, match="MAPPING_NOT_CONFIRMED"):
        reports.create_export(
            template_version_id=str(template["template_version_id"]),
            mapping_version_id=str(mapping["mapping_version_id"]),
            idempotency_key="export-1",
            filters={"factory_id": "FACTORY-A"},
            records=[],
            data_watermark="2026-07-23T00:00:00Z",
            actor_id="finance-1",
        )
    reports.confirm_mapping(str(mapping["mapping_version_id"]), actor_id="finance-1")

    exported = reports.create_export(
        template_version_id=str(template["template_version_id"]),
        mapping_version_id=str(mapping["mapping_version_id"]),
        idempotency_key="export-1",
        filters={"factory_id": "FACTORY-A"},
        records=[{
            "submission_id": "SUB-1",
            "employee_code": "E001",
            "amount": "=2+2",
        }],
        data_watermark="2026-07-23T00:00:00Z",
        actor_id="finance-1",
    )
    downloaded = reports.download(str(exported["export_batch_id"]))
    sheet = load_workbook(BytesIO(downloaded), data_only=False)["工资明细"]
    assert sheet["A2"].value == "E001"
    assert sheet["B2"].value == "'=2+2"
    lineage = reports.lineage(str(exported["export_batch_id"]))["items"]
    assert lineage[0]["submission_id"] == "SUB-1"
    assert {item["cell_address"] for item in lineage} == {"A2", "B2"}
    assert exported["status"] == "AVAILABLE"
    assert exported["data_watermark"] == "2026-07-23T00:00:00Z"
