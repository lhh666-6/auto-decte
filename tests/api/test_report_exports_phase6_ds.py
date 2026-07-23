"""Phase 6 Web report template, mapping, export and lineage flow."""

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.main_ds import create_app
from app.modules.master_data.models_ds import MasterDataCatalog
from app.modules.submission_ledger.service_ds import SubmissionLedgerService
from app.services.container import Services, build_services
from config.settings import Settings


def _add_user(services: Services, code: str, role: str) -> None:
    services.master_data.create(
        MasterDataCatalog.EMPLOYEES, code, code, {}, "phase-6", "fixture"
    )
    services.mobile_identity_repository.set_credential(code, "2468")
    services.mobile_identity_repository.set_access_profile(
        code,
        team_id="",
        team_name="",
        position=role,
        roles=[role],
        allowed_form_types=[],
        allowed_processes=[],
        factory_id="",
        factory_name="",
        bamboo_role="",
    )


@pytest.fixture()
def services(tmp_path: Path) -> Services:
    built = build_services(Settings(data_root=tmp_path))
    _add_user(built, "ADMIN-1", "ADMIN")
    _add_user(built, "FINANCE-1", "FINANCE")
    SubmissionLedgerService(built.engine).record_acceptance(
        submission_id="SUB-1",
        factory_id="FACTORY-A",
        subject_employee_code="E001",
        actor_id="E001",
        definition_version_id="FORM-V1",
        values={"quantity": 9},
        submitted_at=datetime(2026, 7, 23, tzinfo=UTC),
    )
    return built


def _login(services: Services, code: str) -> TestClient:
    client = TestClient(create_app(services))
    assert client.post(
        "/api/v1/web/auth/login",
        json={"employee_code": code, "pin": "2468"},
    ).status_code == 200
    return client


def _csrf(client: TestClient) -> dict[str, str]:
    return {"X-CSRF-Token": client.cookies["web_csrf"]}


def _template() -> bytes:
    workbook = Workbook()
    workbook.active.title = "明细"
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_admin_upload_finance_mapping_export_download_and_lineage(
    services: Services,
) -> None:
    admin = _login(services, "ADMIN-1")
    finance = _login(services, "FINANCE-1")
    uploaded = admin.post(
        "/api/v1/admin/report-templates",
        headers={
            **_csrf(admin),
            "X-Filename": "%E4%BA%A7%E9%87%8F%E6%A8%A1%E6%9D%BF.xlsx",
            "Content-Type": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        },
        content=_template(),
    )
    assert uploaded.status_code == 201, uploaded.text
    template_id = uploaded.json()["template_version_id"]

    mapping = finance.post(
        "/api/v1/finance/report-mappings",
        headers=_csrf(finance),
        json={
            "template_version_id": template_id,
            "mapping_json": {
                "sheet": "明细",
                "start_row": 2,
                "columns": [
                    {"column": 1, "source_field": "subject_employee_code"},
                    {"column": 2, "source_field": "quantity"},
                ],
            },
        },
    ).json()
    mapping_id = mapping["mapping_version_id"]
    finance.post(
        f"/api/v1/finance/report-mappings/{mapping_id}/confirm",
        headers=_csrf(finance),
    )
    exported = finance.post(
        "/api/v1/finance/exports",
        headers=_csrf(finance),
        json={
            "template_version_id": template_id,
            "mapping_version_id": mapping_id,
            "filters": {"factory_id": "FACTORY-A"},
            "idempotency_key": "export-phase-6",
        },
    )
    assert exported.status_code == 201, exported.text
    batch_id = exported.json()["export_batch_id"]
    assert exported.json()["status"] == "AVAILABLE"
    downloaded = finance.get(f"/api/v1/finance/exports/{batch_id}/download")
    assert downloaded.status_code == 200
    sheet = load_workbook(BytesIO(downloaded.content))["明细"]
    assert sheet["A2"].value == "E001"
    assert sheet["B2"].value == 9
    lineage = finance.get(
        f"/api/v1/finance/exports/{batch_id}/lineage"
    ).json()["items"]
    assert {row["cell_address"] for row in lineage} == {"A2", "B2"}
