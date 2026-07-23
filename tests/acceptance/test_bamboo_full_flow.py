"""Multi-role production to finance acceptance flow."""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from .conftest import FACTORY_A, FACTORY_B, AcceptanceEnvironment


def _submit_stage(
    env: AcceptanceEnvironment,
    client,
    *,
    step_id: str,
    actor: str,
    record_id: str,
    stage: str,
    revision: int,
    values: dict[str, object],
    key: str,
) -> dict[str, object]:
    with env.recorder.step(step_id, f"提交 {stage} 环节", actor, "MOBILE") as step:
        response = client.post(
            step,
            f"/api/v1/mobile/bamboo/records/{record_id}/stages/{stage}/submit",
            expected_status=200,
            headers=env.mobile_headers(client, key),
            json={
                "expected_revision": revision,
                "device_id": f"{actor}-{stage}-phone",
                "values": values,
            },
        )
        payload = response.json()
        step.attach(record_id=record_id, stage=stage, revision_after=payload["revision"])
        return payload


def test_multi_role_mobile_to_finance_summary_and_xlsx(
    acceptance_env: AcceptanceEnvironment,
) -> None:
    env = acceptance_env
    identities = (
        ("ACC-SORT", "自动分选工", "SORT_OPERATOR", FACTORY_A, ["WORKER"]),
        ("ACC-DIP", "自动浸胶工", "DIPPING_OPERATOR", FACTORY_A, ["WORKER"]),
        ("ACC-DRY", "自动干燥工", "DRYING_RACK_OPERATOR", FACTORY_A, ["WORKER"]),
        ("ACC-SUP", "自动主管", "SUPERVISOR", FACTORY_A, ["WORKER"]),
        (
            "ACC-PLANT",
            "自动厂长",
            "PLANT_MANAGER",
            FACTORY_A,
            ["PLANT_MANAGER"],
        ),
        (
            "ACC-FIN",
            "自动财务",
            "FINANCE_APPROVER",
            FACTORY_A,
            ["FINANCE"],
        ),
        ("ACC-OTHER", "二厂分选工", "SORT_OPERATOR", FACTORY_B, ["WORKER"]),
    )
    for code, name, bamboo_role, factory_id, workspace_roles in identities:
        env.add_identity(
            employee_code=code,
            employee_name=name,
            bamboo_role=bamboo_role,
            factory_id=factory_id,
            workspace_roles=workspace_roles,
        )

    sort_client = env.mobile("ACC-SORT")
    dipping_client = env.mobile("ACC-DIP")
    drying_client = env.mobile("ACC-DRY")
    supervisor_client = env.mobile("ACC-SUP")
    finance_client = env.mobile("ACC-FIN")
    other_client = env.mobile("ACC-OTHER")
    plant_web = env.web("ACC-PLANT")
    plant_mobile = env.mobile("ACC-PLANT")

    run_id = uuid4().hex[:10].upper()
    with env.recorder.step(
        "01-create-sorting",
        "分选工新建竹丝记录",
        "ACC-SORT",
        "MOBILE",
    ) as step:
        response = sort_client.post(
            step,
            "/api/v1/mobile/bamboo/records",
            expected_status=201,
            headers=env.mobile_headers(sort_client, f"{run_id}-create"),
            json={
                "base_info": {
                    "mode": "分选+装笼",
                    "cage_no": f"ACC-{run_id}",
                    "length": "2.3",
                    "shade": "深",
                    "grade": "A",
                    "bundle_count": 16,
                }
            },
        )
        sorting = response.json()
        sorting_id = str(sorting["record_id"])
        step.attach(record_id=sorting_id, revision=sorting["revision"])

    with env.recorder.step(
        "02-cross-factory-isolation",
        "二厂账号不能读取一厂记录",
        "ACC-OTHER",
        "MOBILE",
    ) as step:
        response = other_client.get(
            step,
            f"/api/v1/mobile/bamboo/records/{sorting_id}",
            expected_status=404,
        )
        step.check(response.json()["code"] == "RECORD_NOT_VISIBLE", "跨厂错误码不稳定")

    sorting = _submit_stage(
        env,
        sort_client,
        step_id="03-submit-sort",
        actor="ACC-SORT",
        record_id=sorting_id,
        stage="SORT",
        revision=int(sorting["revision"]),
        values={"moisture": [12, 13, 14]},
        key=f"{run_id}-sort",
    )
    with env.recorder.step(
        "04-dipping-receives-upstream",
        "浸胶岗位获得上游生成的联合表",
        "ACC-DIP",
        "MOBILE",
    ) as step:
        response = dipping_client.get(
            step,
            "/api/v1/mobile/bamboo/tasks",
            expected_status=200,
        )
        linked = next(
            item
            for item in response.json()["tasks"]
            if item["source_record_id"] == sorting_id
        )
        linked_id = str(linked["record_id"])
        step.attach(record_id=linked_id, source_record_id=sorting_id)
        step.check(linked["form_type"] == "DIPPING_DRYING", "下游表单类型错误")

    linked = _submit_stage(
        env,
        dipping_client,
        step_id="05-submit-dipping",
        actor="ACC-DIP",
        record_id=linked_id,
        stage="DIPPING",
        revision=int(linked["revision"]),
        values={"moisture": [11, 12, 13], "glue_gain": "5"},
        key=f"{run_id}-dipping",
    )
    linked = _submit_stage(
        env,
        drying_client,
        step_id="06-submit-drying",
        actor="ACC-DRY",
        record_id=linked_id,
        stage="DRYING",
        revision=int(linked["revision"]),
        values={
            "moisture": [8, 9, 10],
            "rack_numbers": [f"{run_id}-R01", f"{run_id}-R02"],
        },
        key=f"{run_id}-drying",
    )
    sorting = _submit_stage(
        env,
        supervisor_client,
        step_id="07-supervisor-sorting",
        actor="ACC-SUP",
        record_id=sorting_id,
        stage="SUPERVISOR",
        revision=int(sorting["revision"]),
        values={},
        key=f"{run_id}-sup-sort",
    )
    linked = _submit_stage(
        env,
        supervisor_client,
        step_id="08-supervisor-linked",
        actor="ACC-SUP",
        record_id=linked_id,
        stage="SUPERVISOR",
        revision=int(linked["revision"]),
        values={},
        key=f"{run_id}-sup-linked",
    )

    with env.recorder.step(
        "09-plant-mobile-blocked",
        "厂长移动端生产接口被统一阻止",
        "ACC-PLANT",
        "MOBILE",
    ) as step:
        response = plant_mobile.get(
            step,
            "/api/v1/mobile/bamboo/dashboard",
            expected_status=403,
        )
        step.check(
            response.json()["code"] == "PLANT_MANAGER_WEB_ONLY",
            "厂长移动端边界错误",
        )

    for record_id, label in ((sorting_id, "sort"), (linked_id, "linked")):
        with env.recorder.step(
            f"10-terminate-window-{label}",
            "厂长结束检测窗口",
            "ACC-PLANT",
            "WEB",
        ) as step:
            plant_web.post(
                step,
                f"/api/v1/plant/inspection-queue/{record_id}/terminate",
                expected_status=200,
                headers=env.web_headers(plant_web, f"{run_id}-terminate-{label}"),
                json={"confirm": True},
            )

    audited: dict[str, dict[str, object]] = {}
    for record, label in ((sorting, "sort"), (linked, "linked")):
        record_id = str(record["record_id"])
        with env.recorder.step(
            f"11-plant-audit-{label}",
            "厂长 Web 审核生产记录",
            "ACC-PLANT",
            "WEB",
        ) as step:
            response = plant_web.post(
                step,
                f"/api/v1/plant/records/{record_id}/audit",
                expected_status=200,
                headers=env.web_headers(plant_web, f"{run_id}-audit-{label}"),
                json={
                    "expected_revision": record["revision"],
                    "device_id": "acceptance-plant-web",
                    "values": {},
                },
            )
            audited[label] = response.json()
            step.check(audited[label]["status"] == "COMPLETED", "厂长审核后未完成")

    with env.recorder.step(
        "12-finance-daily-batch",
        "财务读取日批次并验证工资事实",
        "ACC-FIN",
        "MOBILE",
    ) as step:
        response = finance_client.get(
            step,
            "/api/v1/mobile/bamboo/finance/daily-batches",
            expected_status=200,
        )
        batches = response.json()
        batch = next(
            item
            for item in batches
            if {entry["record_id"] for entry in item["items"]}
            >= {sorting_id, linked_id}
        )
        relevant = [
            item
            for item in batch["items"]
            if item["record_id"] in {sorting_id, linked_id}
        ]
        expected_amounts = {
            "ACC-SORT": "96.00",
            "ACC-DIP": "5.00",
            "ACC-DRY": "2.00",
        }
        step.check(len(relevant) == 3, "工资条目数量不是3")
        step.check(
            {item["employee_code"]: item["amount"] for item in relevant}
            == expected_amounts,
            "工资计算结果不正确",
        )
        step.attach(batch_id=batch["batch_id"], business_date=batch["business_date"])

    for item in relevant:
        with env.recorder.step(
            f"13-finance-approve-{item['employee_code']}",
            "财务审批工资条目",
            "ACC-FIN",
            "MOBILE",
        ) as step:
            response = finance_client.post(
                step,
                f"/api/v1/mobile/bamboo/finance/items/{item['item_id']}/decision",
                expected_status=200,
                headers=env.mobile_headers(
                    finance_client,
                    f"{run_id}-approve-{item['item_id']}",
                ),
                json={"decision": "APPROVED", "note": "自动验收通过"},
            )
            step.check(response.json()["status"] == "APPROVED", "财务审批状态错误")

    month = str(batch["business_date"])[:7]
    with env.recorder.step(
        "14-monthly-summary",
        "验证财务月汇总",
        "ACC-FIN",
        "MOBILE",
    ) as step:
        response = finance_client.get(
            step,
            "/api/v1/mobile/bamboo/finance/monthly-summary",
            expected_status=200,
            params={"month": month},
        )
        summary = response.json()
        step.check(summary["total_amount"] == "103.00", "月汇总总额应为103.00")
        step.check(
            {item["employee_code"]: item["amount"] for item in summary["items"]}
            == expected_amounts,
            "月汇总员工金额错误",
        )

    with env.recorder.step(
        "15-xlsx-export",
        "下载并验证财务 XLSX",
        "ACC-FIN",
        "MOBILE",
    ) as step:
        response = finance_client.get(
            step,
            "/api/v1/mobile/bamboo/finance/export.xlsx",
            expected_status=200,
            params={"month": month},
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        try:
            sheet = workbook["工资审计"]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        step.check(rows[0] == ("工号", "已审批工资"), "导出表头错误")
        step.check(
            {str(code): str(amount) for code, amount in rows[1:]} == expected_amounts,
            "导出员工金额错误",
        )
