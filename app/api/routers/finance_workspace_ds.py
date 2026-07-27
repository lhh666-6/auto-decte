"""Phase 1 finance workspace endpoints."""

from urllib.parse import quote

from fastapi import APIRouter, Header, HTTPException, Query, Request, status
from fastapi.responses import Response

from app.api.routers.web_auth_ds import require_web_actor, require_web_csrf
from app.api.schemas.business_workflows_ds import (
    ConfirmDiscoveryRequest,
    CreateDiscoverySessionRequest,
    CreateWorkflowRequest,
    DiscoveryMessageRequest,
)
from app.api.schemas.managed_forms_ds import (
    CreateManagedFormRequest,
    ManagedFormListResponse,
    ManagedFormVersionResponse,
    UpdateManagedFormVersionRequest,
)
from app.api.schemas.payroll_rules_ds import (
    CalculatePayrollRequest,
    CreatePayrollRuleRequest,
)
from app.api.schemas.report_templates_ds import (
    CreateGovernedExportRequest,
    CreateReportMappingRequest,
    PreviewGovernedExportRequest,
)
from app.api.schemas.submission_ledger_ds import (
    AttachReplacementRequest,
    CreateCorrectionRequest,
    ReviewCorrectionRequest,
)
from app.api.schemas.web_workspaces_ds import OverviewCard, WorkspaceOverviewResponse
from app.modules.business_discovery.service_ds import (
    BusinessDiscoveryError,
    BusinessDiscoveryService,
)
from app.modules.electronic_forms.governance_ds import ManagedFormError, ManagedFormService
from app.modules.identity_access.web_policy_ds import WebWorkspace, allows_workspace
from app.modules.payroll_rules.service_ds import PayrollError, PayrollService
from app.modules.report_templates.service_ds import (
    ReportTemplateError,
    ReportTemplateService,
)
from app.modules.submission_ledger.service_ds import (
    SubmissionLedgerError,
    SubmissionLedgerService,
)
from app.modules.workflow_engine.service_ds import WorkflowError, WorkflowService

router = APIRouter(prefix="/api/v1/finance", tags=["finance"])


def _finance_actor(request: Request):  # type: ignore[no-untyped-def]
    actor = require_web_actor(request)
    if not allows_workspace(actor, WebWorkspace.FINANCE):
        raise HTTPException(
            status_code=403,
            detail={"code": "WEB_ROLE_FORBIDDEN", "detail": "当前账号不能进入财务工作区。"},
        )
    return actor


def _forms(request: Request) -> ManagedFormService:
    return ManagedFormService(request.app.state.services.engine)


def _discovery(request: Request) -> BusinessDiscoveryService:
    return BusinessDiscoveryService(request.app.state.services.engine)


def _workflows(request: Request) -> WorkflowService:
    return WorkflowService(request.app.state.services.engine)


def _ledger(request: Request) -> SubmissionLedgerService:
    return SubmissionLedgerService(request.app.state.services.engine)


def _payroll(request: Request) -> PayrollService:
    return PayrollService(request.app.state.services.engine)


def _reports(request: Request) -> ReportTemplateService:
    return ReportTemplateService(request.app.state.services.engine)


def _form_error(error: ManagedFormError) -> HTTPException:
    conflict_codes = {
        "FORM_KEY_EXISTS",
        "FORM_VERSION_IMMUTABLE",
        "FORM_VERSION_REVISION_CONFLICT",
        "FORM_VERSION_NOT_DRAFT",
    }
    return HTTPException(
        status_code=409 if error.code in conflict_codes else 404,
        detail={"code": error.code, "detail": error.detail},
    )


@router.get("/overview", response_model=WorkspaceOverviewResponse)
def overview(
    request: Request,
    factory_id: str = Query(default=""),
) -> WorkspaceOverviewResponse:
    actor = require_web_actor(request)
    if not allows_workspace(actor, WebWorkspace.FINANCE):
        raise HTTPException(
            status_code=403,
            detail={"code": "WEB_ROLE_FORBIDDEN", "detail": "当前账号不能进入财务工作区。"},
        )
    return WorkspaceOverviewResponse(
        workspace=WebWorkspace.FINANCE.value,
        title="财务概览",
        scope="GLOBAL",
        factory_id=factory_id,
        cards=[
            OverviewCard(key="today_submissions", label="今日新增提交", value=0),
            OverviewCard(key="pending_review", label="待审核记录", value=0),
            OverviewCard(key="exceptions", label="异常记录", value=0),
        ],
    )


@router.get("/ledger/overview")
def ledger_overview(
    request: Request,
    factory_id: str = Query(default=""),
    scope: str = Query(default=""),
) -> dict[str, int]:
    _finance_actor(request)
    return _ledger(request).overview(factory_id or None, scope or None)


@router.get("/ledger")
def ledger(
    request: Request,
    factory_id: str = Query(default=""),
    scope: str = Query(default=""),
) -> dict[str, object]:
    _finance_actor(request)
    return _ledger(request).list_ledger(factory_id or None, scope or None)


@router.post(
    "/ledger/{submission_id}/corrections",
    status_code=status.HTTP_201_CREATED,
)
def create_correction(
    submission_id: str,
    body: CreateCorrectionRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key"),
) -> dict[str, str]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    if idempotency_key is None or not idempotency_key.strip():
        raise HTTPException(
            status_code=400,
            detail={
                "code": "IDEMPOTENCY_KEY_REQUIRED",
                "detail": "更正请求必须提供 Idempotency-Key。",
            },
        )
    ledger_svc = _ledger(request)
    try:
        effective = ledger_svc.get_effective(submission_id)
    except SubmissionLedgerError as error:
        raise HTTPException(
            status_code=404,
            detail={"code": error.code, "detail": error.detail},
        ) from error
    factory_id = str(effective["factory_id"])
    assigned_to = str(effective["subject_employee_code"])
    try:
        return ledger_svc.return_submission(
            submission_id=submission_id,
            factory_id=factory_id,
            reason=body.reason,
            requested_by=actor.employee_code,
            assigned_to=assigned_to,
            idempotency_key=idempotency_key.strip(),
            correction_type=body.correction_type,
            supplementary_note=body.supplementary_note,
        )
    except SubmissionLedgerError as error:
        status_code_http = (
            409
            if error.code != "IDEMPOTENCY_CONFLICT"
            else status.HTTP_409_CONFLICT
        )
        raise HTTPException(
            status_code=status_code_http,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/corrections")
def corrections(request: Request, factory_id: str = Query(default="")) -> dict[str, object]:
    _finance_actor(request)
    return _ledger(request).list_corrections(factory_id or None)


@router.get("/business-tasks")
def business_tasks(request: Request, factory_id: str = Query(default="")) -> dict[str, object]:
    _finance_actor(request)
    return _ledger(request).list_tasks(factory_id or None)


@router.get("/payroll-rules")
def payroll_rules(request: Request) -> dict[str, object]:
    _finance_actor(request)
    return _payroll(request).list_rules()


@router.post("/payroll-rules", status_code=status.HTTP_201_CREATED)
def create_payroll_rule(
    body: CreatePayrollRuleRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _payroll(request).create_rule(
            **body.model_dump(), actor_id=actor.employee_code
        )
    except PayrollError as error:
        raise HTTPException(
            status_code=422,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/payroll-rules/{version_id}/submit-approval")
def submit_payroll_rule(
    version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _payroll(request).submit_rule(version_id)
    except PayrollError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/payroll-calculations", status_code=status.HTTP_201_CREATED)
def calculate_payroll(
    body: CalculatePayrollRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _payroll(request).calculate(
            **body.model_dump(), actor_id=actor.employee_code
        )
    except PayrollError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/payroll-calculations")
def payroll_batches(request: Request) -> dict[str, object]:
    _finance_actor(request)
    return _payroll(request).list_batches()


@router.post("/payroll-calculations/{batch_id}/confirm")
def confirm_payroll_batch(
    batch_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _payroll(request).confirm_batch(batch_id, actor_id=actor.employee_code)
    except PayrollError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/payroll")
def finance_payroll(
    request: Request,
    factory_id: str = Query(default=""),
    employee_code: str = Query(default=""),
) -> dict[str, object]:
    actor = _finance_actor(request)
    return _payroll(request).list_official(
        factory_id=factory_id or None,
        employee_code=employee_code or None,
        actor_id=actor.employee_code,
        actor_role="FINANCE",
    )


@router.get("/report-templates")
def report_templates(request: Request) -> dict[str, object]:
    _finance_actor(request)
    return _reports(request).list_templates()


@router.post("/report-templates/{template_version_id}/analyze")
def analyze_report_template(
    template_version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _reports(request).get_template(template_version_id)
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=404,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/report-mappings")
def report_mappings(
    request: Request,
    template_version_id: str = Query(default=""),
) -> dict[str, object]:
    _finance_actor(request)
    return _reports(request).list_mappings(template_version_id or None)


@router.post("/report-mappings", status_code=201)
def create_report_mapping(
    body: CreateReportMappingRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _reports(request).create_mapping(
            **body.model_dump(), actor_id=actor.employee_code
        )
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=422,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/report-mappings/{mapping_version_id}/confirm")
def confirm_report_mapping(
    mapping_version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _reports(request).confirm_mapping(
            mapping_version_id, actor_id=actor.employee_code
        )
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/exports/preview")
def preview_governed_export(
    body: PreviewGovernedExportRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    ledger_service = _ledger(request)
    factory_id = str(body.filters.get("factory_id", "")).strip() or None
    source_records = ledger_service.list_ledger(factory_id)["items"]
    records: list[dict[str, object]] = []
    if not isinstance(source_records, list):
        source_records = []
    for item in source_records:
        if not isinstance(item, dict):
            continue
        record = dict(item)
        values = record.pop("values", {})
        if isinstance(values, dict):
            record.update(values)
        record["submission_id"] = record["effective_submission_id"]
        records.append(record)
    try:
        return _reports(request).preview_export(
            template_version_id=body.template_version_id,
            mapping_version_id=body.mapping_version_id,
            filters=body.filters,
            records=records,
        )
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/exports", status_code=201)
def create_governed_export(
    body: CreateGovernedExportRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    ledger_service = _ledger(request)
    factory_id = str(body.filters.get("factory_id", "")).strip() or None
    source_records = ledger_service.list_ledger(factory_id)["items"]
    records: list[dict[str, object]] = []
    if not isinstance(source_records, list):
        source_records = []
    for item in source_records:
        if not isinstance(item, dict):
            continue
        record = dict(item)
        values = record.pop("values", {})
        if isinstance(values, dict):
            record.update(values)
        record["submission_id"] = record["effective_submission_id"]
        records.append(record)
    try:
        return _reports(request).create_export(
            template_version_id=body.template_version_id,
            mapping_version_id=body.mapping_version_id,
            idempotency_key=body.idempotency_key,
            filters=body.filters,
            records=records,
            data_watermark=ledger_service.watermark(),
            actor_id=actor.employee_code,
        )
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/exports/{source_batch_id}/reexport", status_code=201)
def reexport_governed_export(
    source_batch_id: str,
    body: CreateGovernedExportRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    ledger_service = _ledger(request)
    factory_id = str(body.filters.get("factory_id", "")).strip() or None
    source_records = ledger_service.list_ledger(factory_id)["items"]
    records: list[dict[str, object]] = []
    if not isinstance(source_records, list):
        source_records = []
    for item in source_records:
        if not isinstance(item, dict):
            continue
        record = dict(item)
        values = record.pop("values", {})
        if isinstance(values, dict):
            record.update(values)
        record["submission_id"] = record["effective_submission_id"]
        records.append(record)
    try:
        return _reports(request).reexport(
            source_batch_id=source_batch_id,
            idempotency_key=body.idempotency_key,
            records=records,
            data_watermark=ledger_service.watermark(),
            actor_id=actor.employee_code,
        )
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/exports")
def governed_exports(request: Request) -> dict[str, object]:
    _finance_actor(request)
    return _reports(request).list_exports()


@router.get("/exports/{export_batch_id}")
def governed_export(export_batch_id: str, request: Request) -> dict[str, object]:
    _finance_actor(request)
    try:
        return _reports(request).get_export(export_batch_id)
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=404,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/exports/{export_batch_id}/download")
def download_governed_export(export_batch_id: str, request: Request) -> Response:
    _finance_actor(request)
    try:
        batch = _reports(request).get_export(export_batch_id)
        content = _reports(request).download(export_batch_id)
    except ReportTemplateError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f"attachment; filename*=UTF-8''{quote(str(batch['download_name']))}"
            )
        },
    )


@router.get("/exports/{export_batch_id}/lineage")
def governed_export_lineage(
    export_batch_id: str, request: Request
) -> dict[str, object]:
    _finance_actor(request)
    return _reports(request).lineage(export_batch_id)


@router.post("/corrections/{correction_id}/replacement")
def attach_correction_replacement(
    correction_id: str,
    body: AttachReplacementRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, str]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _ledger(request).attach_replacement(correction_id, **body.model_dump())
    except SubmissionLedgerError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/corrections/{correction_id}/review")
def review_correction(
    correction_id: str,
    body: ReviewCorrectionRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, str]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _ledger(request).review_correction(
            correction_id,
            approved=body.approved,
            reviewer_id=actor.employee_code,
            note=body.note,
        )
    except SubmissionLedgerError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.get("/form-definitions", response_model=ManagedFormListResponse)
def form_definitions(request: Request) -> ManagedFormListResponse:
    _finance_actor(request)
    return ManagedFormListResponse(
        items=[
            ManagedFormVersionResponse.model_validate(item)
            for item in _forms(request).list_definitions()
        ]
    )


@router.post(
    "/form-definitions",
    response_model=ManagedFormVersionResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_form_definition(
    body: CreateManagedFormRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> ManagedFormVersionResponse:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        payload = _forms(request).create_definition(
            **body.model_dump(by_alias=True),
            actor_id=actor.employee_code,
        )
    except ManagedFormError as error:
        raise _form_error(error) from error
    return ManagedFormVersionResponse.model_validate(payload)


@router.get(
    "/form-versions/{version_id}",
    response_model=ManagedFormVersionResponse,
)
def form_version(version_id: str, request: Request) -> ManagedFormVersionResponse:
    _finance_actor(request)
    try:
        payload = _forms(request).get_version(version_id)
    except ManagedFormError as error:
        raise _form_error(error) from error
    return ManagedFormVersionResponse.model_validate(payload)


@router.patch(
    "/form-versions/{version_id}",
    response_model=ManagedFormVersionResponse,
)
def update_form_version(
    version_id: str,
    body: UpdateManagedFormVersionRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> ManagedFormVersionResponse:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        payload = _forms(request).update_draft(
            version_id,
            **body.model_dump(by_alias=True),
        )
    except ManagedFormError as error:
        raise _form_error(error) from error
    return ManagedFormVersionResponse.model_validate(payload)


@router.post(
    "/form-versions/{version_id}/submit-approval",
    response_model=ManagedFormVersionResponse,
)
def submit_form_approval(
    version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> ManagedFormVersionResponse:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        payload = _forms(request).submit_approval(version_id)
    except ManagedFormError as error:
        raise _form_error(error) from error
    return ManagedFormVersionResponse.model_validate(payload)


@router.post(
    "/form-versions/{version_id}/clone",
    response_model=ManagedFormVersionResponse,
    status_code=status.HTTP_201_CREATED,
)
def clone_form_version(
    version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> ManagedFormVersionResponse:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        payload = _forms(request).clone(version_id, actor.employee_code)
    except ManagedFormError as error:
        raise _form_error(error) from error
    return ManagedFormVersionResponse.model_validate(payload)


@router.post("/business-discovery/sessions", status_code=status.HTTP_201_CREATED)
def create_discovery_session(
    body: CreateDiscoverySessionRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    return _discovery(request).create_session(
        **body.model_dump(),
        actor_id=actor.employee_code,
    )


@router.post("/business-discovery/sessions/{session_id}/messages")
def add_discovery_message(
    session_id: str,
    body: DiscoveryMessageRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _discovery(request).add_message(session_id, content=body.content)
    except BusinessDiscoveryError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/business-discovery/sessions/{session_id}/confirm")
def confirm_discovery(
    session_id: str,
    body: ConfirmDiscoveryRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _discovery(request).confirm(
            session_id,
            **body.model_dump(),
            actor_id=actor.employee_code,
        )
    except BusinessDiscoveryError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/workflows", status_code=status.HTTP_201_CREATED)
def create_workflow(
    body: CreateWorkflowRequest,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    actor = _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _workflows(request).create(
            **body.model_dump(),
            actor_id=actor.employee_code,
        )
    except WorkflowError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/workflow-versions/{version_id}/validate")
def validate_workflow(
    version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _workflows(request).validate(version_id)
    except WorkflowError as error:
        raise HTTPException(
            status_code=404,
            detail={"code": error.code, "detail": error.detail},
        ) from error


@router.post("/workflow-versions/{version_id}/submit-approval")
def submit_workflow(
    version_id: str,
    request: Request,
    x_csrf_token: str | None = Header(default=None, alias="X-CSRF-Token"),
) -> dict[str, object]:
    _finance_actor(request)
    require_web_csrf(request, x_csrf_token)
    try:
        return _workflows(request).submit(version_id)
    except WorkflowError as error:
        raise HTTPException(
            status_code=409,
            detail={"code": error.code, "detail": error.detail},
        ) from error


# ─────────────────────────────────────────────────────────────
# V1 Payroll Field Registry
# ─────────────────────────────────────────────────────────────


@router.get("/payroll-fields")
def payroll_fields(
    request: Request,
    position: str = Query(default=""),
) -> dict[str, object]:
    """Return fields registered for a position, with usage classification."""
    _finance_actor(request)
    from app.adapters.database.models import PayrollFieldRegistryRow
    from sqlalchemy import select as sa_select
    from sqlalchemy.orm import Session as SaSession

    engine = request.app.state.services.engine
    with SaSession(engine) as session:
        statement = sa_select(PayrollFieldRegistryRow).where(
            PayrollFieldRegistryRow.active.is_(True),
        )
        if position.strip():
            statement = statement.where(
                PayrollFieldRegistryRow.position_role == position.strip().upper(),
            )
        rows = session.scalars(statement).all()
        items: list[dict[str, str]] = []
        for row in rows:
            usage = "LOOKUP_DIMENSION" if row.data_type == "STRING" else "NUMERIC_METRIC"
            items.append({
                "field_key": row.field_key,
                "display_name": row.display_name,
                "data_type": row.data_type,
                "usage": usage,
            })
    return {"items": items}


# ─────────────────────────────────────────────────────────────
# V1 Position Data (simplified finance)
# ─────────────────────────────────────────────────────────────

from typing import Any, cast  # noqa: E402

from app.application.bamboo_operations_ds import BambooOperationsService  # noqa: E402


def _bamboo(request: Request) -> BambooOperationsService:
    return cast(BambooOperationsService, request.app.state.services.bamboo_operations)


# ── V1 Runtime Closure §12: Finance factories endpoint ──


@router.get("/factories")
def finance_factories(request: Request) -> dict[str, object]:
    """Finance read-only factory list. Returns { items: [...] }."""
    _finance_actor(request)
    factories = _bamboo(request).list_factories_for_admin()
    return {"items": factories}


# ── V1 Runtime Closure §19.3: Finance Management Salary read-only ──


@router.get("/management-salaries")
def finance_management_salaries(
    request: Request,
    factory_id: str | None = None,
) -> dict[str, object]:
    """Finance read-only view of management salaries."""
    _finance_actor(request)
    services: Any = request.app.state.services
    items = services.management_salary.list_salaries(factory_id)
    return {"items": items}


@router.get("/position-data")
def position_data(
    request: Request,
    factory_id: str = Query(default=""),
    stage: str = Query(default=""),
    date_from: str = Query(default=""),
    date_to: str = Query(default=""),
    employee_code: str = Query(default=""),
) -> dict[str, object]:
    """V1: List production/payroll data filtered by position (stage)."""
    _finance_actor(request)
    items = _bamboo(request).list_position_data(
        factory_id=factory_id or None,
        stage=stage or None,
        date_from=date_from or None,
        date_to=date_to or None,
        employee_code=employee_code or None,
    )
    return {"items": items, "count": len(items)}


@router.get("/position-data/export")
def export_position_data(
    request: Request,
    factory_id: str = Query(default=""),
    stage: str = Query(default=""),
    date_from: str = Query(default=""),
    date_to: str = Query(default=""),
    employee_code: str = Query(default=""),
) -> Response:
    """V1: Export position data as XLSX with fixed column schemas per position.

    V1 Runtime Closure §13: Fixed columns per stage, NOT dynamic sample.values().keys().
    Internal IDs (record_id, submission_id) excluded from main sheet.
    """
    _finance_actor(request)
    from io import BytesIO
    from datetime import UTC, datetime

    items = _bamboo(request).list_position_data(
        factory_id=factory_id or None,
        stage=stage or None,
        date_from=date_from or None,
        date_to=date_to or None,
        employee_code=employee_code or None,
    )

    # V1: Fixed column schemas per position
    SORT_COLUMNS = [
        "日期", "员工工号", "员工姓名", "工厂", "表号", "笼号",
        "把数", "长度", "深浅", "原评级", "最终评级", "净重", "含水率", "状态",
    ]
    DIPPING_COLUMNS = [
        "日期", "员工工号", "员工姓名", "工厂", "表号", "笼号",
        "胶前重", "胶后重", "上胶量", "胶液批次", "浸胶开始", "浸胶结束",
        "含水率", "原评级", "最终评级", "状态",
    ]
    DRYING_COLUMNS = [
        "日期", "员工工号", "员工姓名", "工厂", "表号", "笼号",
        "干燥架号", "架数", "干燥开始", "干燥结束",
        "含水率", "原评级", "最终评级", "状态",
    ]

    # Value-key to column label mapping per stage
    def _val(vals: dict, key: str) -> str:
        raw = vals.get(key, "")
        if isinstance(raw, list):
            return ", ".join(str(x) for x in raw)
        return str(raw) if raw else ""

    SORT_VALUE_KEYS = {
        "bundle_count": "把数", "length": "长度", "shade": "深浅",
        "net_weight": "净重", "moisture": "含水率",
    }
    DIPPING_VALUE_KEYS = {
        "glue_before_weight": "胶前重", "glue_after_weight": "胶后重",
        "glue_gain": "上胶量", "glue_batch": "胶液批次",
        "dipping_start": "浸胶开始", "dipping_end": "浸胶结束",
        "moisture": "含水率",
    }
    DRYING_VALUE_KEYS = {
        "rack_numbers": "干燥架号", "rack_count": "架数",
        "drying_start": "干燥开始", "drying_end": "干燥结束",
        "moisture": "含水率",
    }

    def _make_row(item: dict, columns: list[str], value_keys: dict[str, str]) -> list[str]:
        vals = item.get("values", {}) or {}
        row_data = {
            "日期": (item.get("date", "") or "")[:10],
            "员工工号": item.get("employee_code", ""),
            "员工姓名": item.get("employee_name", ""),
            "工厂": item.get("factory_id", ""),
            "表号": item.get("display_no", ""),
            "笼号": item.get("cage_no", ""),
            "原评级": item.get("original_grade", ""),
            "最终评级": item.get("effective_grade", ""),
            "状态": item.get("status", ""),
        }
        for vk, label in value_keys.items():
            row_data[label] = _val(vals, vk)
        return [row_data.get(col, "") for col in columns]

    stage_col = stage or ""

    from openpyxl import Workbook
    wb = Workbook()

    if stage_col in ("SORT", "DIPPING", "DRYING"):
        if stage_col == "SORT":
            columns, value_keys = SORT_COLUMNS, SORT_VALUE_KEYS
        elif stage_col == "DIPPING":
            columns, value_keys = DIPPING_COLUMNS, DIPPING_VALUE_KEYS
        else:
            columns, value_keys = DRYING_COLUMNS, DRYING_VALUE_KEYS
        ws = wb.active
        ws.title = "岗位数据"
        ws.append(columns)
        for item in items:
            ws.append(_make_row(item, columns, value_keys))
    else:
        # "全部岗位": three sheets — SORT / DIPPING / DRYING
        stage_specs = [
            ("分选工", SORT_COLUMNS, SORT_VALUE_KEYS, "SORT"),
            ("浸胶工", DIPPING_COLUMNS, DIPPING_VALUE_KEYS, "DIPPING"),
            ("干燥工", DRYING_COLUMNS, DRYING_VALUE_KEYS, "DRYING"),
        ]
        for idx, (sheet_title, columns, value_keys, stage_filter) in enumerate(stage_specs):
            if idx == 0:
                ws = wb.active
            else:
                ws = wb.create_sheet()
            ws.title = sheet_title
            ws.append(columns)
            stage_items = [i for i in items if i.get("stage") == stage_filter]
            for item in stage_items:
                ws.append(_make_row(item, columns, value_keys))

    # V1 §13: Add metadata sheet
    ws_meta = wb.create_sheet("导出说明")
    ws_meta.append(["字段", "值"])
    ws_meta.append(["导出时间", datetime.now(UTC).isoformat()])
    ws_meta.append(["工厂", factory_id or "全部"])
    ws_meta.append(["工序", stage_col or "全部"])
    ws_meta.append(["日期范围", f"{date_from or '不限'} ~ {date_to or '不限'}"])
    ws_meta.append(["记录数", str(len(items))])
    ws_meta.append(["schema_version", "V1_FIXED_SCHEMA"])

    dest = BytesIO()
    wb.save(dest)
    dest.seek(0)
    stage_label = stage or "all"
    filename = f"position-data-{stage_label}.xlsx"
    return Response(
        content=dest.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
