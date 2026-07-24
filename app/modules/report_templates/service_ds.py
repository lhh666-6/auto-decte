"""Safe XLS/XLSX ingestion, governed mappings, exports and cell lineage."""

from __future__ import annotations

import hashlib
import json
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import xlrd  # type: ignore[import-untyped]
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import Engine, func, select
from sqlalchemy.orm import Session

from app.adapters.database.models import (
    ExportCellLineageRow,
    GovernedExportBatchRow,
    ReportMappingVersionRow,
    ReportTemplateVersionRow,
)

MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


class ReportTemplateError(ValueError):
    def __init__(self, code: str, detail: str) -> None:
        super().__init__(f"{code}: {detail}")
        self.code = code
        self.detail = detail


class ReportTemplateService:
    def __init__(
        self,
        engine: Engine,
        clock: Callable[[], datetime] | None = None,
    ) -> None:
        self._engine = engine
        self._clock = clock or (lambda: datetime.now(UTC))

    def upload_template(
        self,
        *,
        filename: str,
        mime_type: str,
        content: bytes,
        actor_id: str,
    ) -> dict[str, object]:
        if not content or len(content) > MAX_TEMPLATE_BYTES:
            raise ReportTemplateError("FILE_SIZE_INVALID", "模板为空或超过 10MB。")
        suffix = Path(filename).suffix.lower()
        if suffix not in {".xls", ".xlsx"}:
            raise ReportTemplateError("FILE_TYPE_INVALID", "只允许 .xls 或 .xlsx。")
        if suffix == ".xlsx":
            if not content.startswith(b"PK"):
                raise ReportTemplateError("FILE_SIGNATURE_INVALID", "XLSX 文件签名无效。")
            self._check_zip(content)
            structure, warnings = self._analyze_xlsx(content)
            format_name = "XLSX"
        else:
            if not content.startswith(OLE_SIGNATURE):
                raise ReportTemplateError("FILE_SIGNATURE_INVALID", "XLS 文件签名无效。")
            structure, warnings = self._analyze_xls(content)
            format_name = "XLS"
        file_hash = self._sha(content)
        structure_hash = self._hash_json(structure)
        now = self._clock()
        with Session(self._engine) as session, session.begin():
            existing = session.scalar(
                select(ReportTemplateVersionRow).where(
                    ReportTemplateVersionRow.file_hash == file_hash
                )
            )
            if existing is not None:
                return self._template(existing)
            row = ReportTemplateVersionRow(
                template_version_id=self._id("RTV"),
                filename=Path(filename).name,
                format=format_name,
                mime_type=mime_type,
                file_hash=file_hash,
                size_bytes=len(content),
                content=content,
                structure=structure,
                structure_hash=structure_hash,
                warnings=warnings,
                status="ANALYZED",
                created_by=actor_id,
                created_at=now,
            )
            session.add(row)
            payload = self._template(row)
        return payload

    def list_templates(self) -> dict[str, object]:
        with Session(self._engine) as session:
            rows = session.scalars(
                select(ReportTemplateVersionRow).order_by(
                    ReportTemplateVersionRow.created_at.desc()
                )
            )
            return {"items": [self._template(row) for row in rows]}

    def get_template(self, template_version_id: str) -> dict[str, object]:
        with Session(self._engine) as session:
            row = session.get(ReportTemplateVersionRow, template_version_id)
            if row is None:
                raise ReportTemplateError("TEMPLATE_NOT_FOUND", "模板版本不存在。")
            return self._template(row)

    def create_mapping(
        self,
        *,
        template_version_id: str,
        mapping_json: dict[str, Any],
        actor_id: str,
    ) -> dict[str, object]:
        now = self._clock()
        with Session(self._engine) as session, session.begin():
            template = session.get(ReportTemplateVersionRow, template_version_id)
            if template is None:
                raise ReportTemplateError("TEMPLATE_NOT_FOUND", "模板版本不存在。")
            normalized = self._validate_mapping(template.structure, mapping_json)
            version = int(
                session.scalar(
                    select(func.max(ReportMappingVersionRow.version)).where(
                        ReportMappingVersionRow.template_version_id
                        == template_version_id
                    )
                )
                or 0
            ) + 1
            row = ReportMappingVersionRow(
                mapping_version_id=self._id("RMV"),
                template_version_id=template_version_id,
                version=version,
                mapping_json=normalized,
                content_hash=self._hash_json(normalized),
                status="DRAFT",
                created_by=actor_id,
                created_at=now,
            )
            session.add(row)
            payload = self._mapping(row)
        return payload

    def confirm_mapping(
        self, mapping_version_id: str, *, actor_id: str
    ) -> dict[str, object]:
        now = self._clock()
        with Session(self._engine) as session, session.begin():
            row = session.get(ReportMappingVersionRow, mapping_version_id)
            if row is None or row.status != "DRAFT":
                raise ReportTemplateError("MAPPING_STATE_INVALID", "映射当前不可确认。")
            row.status = "CONFIRMED"
            row.confirmed_by = actor_id
            row.confirmed_at = now
            payload = self._mapping(row)
        return payload

    def list_mappings(self, template_version_id: str | None = None) -> dict[str, object]:
        with Session(self._engine) as session:
            statement = select(ReportMappingVersionRow).order_by(
                ReportMappingVersionRow.created_at.desc()
            )
            if template_version_id:
                statement = statement.where(
                    ReportMappingVersionRow.template_version_id == template_version_id
                )
            return {"items": [self._mapping(row) for row in session.scalars(statement)]}

    def create_export(
        self,
        *,
        template_version_id: str,
        mapping_version_id: str,
        idempotency_key: str,
        filters: dict[str, Any],
        records: list[dict[str, Any]],
        data_watermark: str,
        actor_id: str,
    ) -> dict[str, object]:
        request_payload = {
            "template_version_id": template_version_id,
            "mapping_version_id": mapping_version_id,
            "filters": filters,
            "records": records,
            "data_watermark": data_watermark,
        }
        request_hash = self._hash_json(request_payload)
        now = self._clock()
        with Session(self._engine) as session, session.begin():
            existing = session.scalar(
                select(GovernedExportBatchRow).where(
                    GovernedExportBatchRow.idempotency_key == idempotency_key
                )
            )
            if existing is not None:
                if existing.request_hash != request_hash:
                    raise ReportTemplateError(
                        "IDEMPOTENCY_CONFLICT", "幂等键对应不同导出请求。"
                    )
                return self._batch(existing)
            template = session.get(ReportTemplateVersionRow, template_version_id)
            mapping = session.get(ReportMappingVersionRow, mapping_version_id)
            if template is None or mapping is None:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "模板或映射不存在。")
            if mapping.template_version_id != template_version_id:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "映射不属于所选模板。")
            if mapping.status != "CONFIRMED":
                raise ReportTemplateError(
                    "MAPPING_NOT_CONFIRMED", "未确认映射不能正式导出。"
                )
            workbook = self._workbook(template)
            mapping_data = mapping.mapping_json
            sheet = workbook[str(mapping_data["sheet"])]
            start_row = int(mapping_data["start_row"])
            export_batch_id = self._id("GEB")
            lineage_rows: list[dict[str, object]] = []
            for offset, record in enumerate(records):
                row_index = start_row + offset
                submission_id = str(record.get("submission_id", ""))
                for column in mapping_data["columns"]:
                    column_index = int(column["column"])
                    field_key = str(column["source_field"])
                    value = self._safe_cell(record.get(field_key))
                    cell = sheet.cell(row=row_index, column=column_index)
                    cell.value = value
                    lineage_rows.append(
                        {
                            "sheet_name": sheet.title,
                            "row_index": row_index,
                            "column_index": column_index,
                            "cell_address": cell.coordinate,
                            "submission_id": submission_id,
                            "field_key": field_key,
                        }
                    )
            output = BytesIO()
            workbook.save(output)
            file_content = output.getvalue()
            load_workbook(BytesIO(file_content), read_only=True, data_only=False).close()
            row = GovernedExportBatchRow(
                export_batch_id=export_batch_id,
                idempotency_key=idempotency_key,
                request_hash=request_hash,
                template_version_id=template_version_id,
                mapping_version_id=mapping_version_id,
                filters=filters,
                data_watermark=data_watermark,
                status="AVAILABLE",
                file_content=file_content,
                file_hash=self._sha(file_content),
                download_name=f"{Path(template.filename).stem}-{now:%Y%m%d%H%M%S}.xlsx",
                created_by=actor_id,
                created_at=now,
            )
            session.add(row)
            session.flush()
            for lineage in lineage_rows:
                session.add(
                    ExportCellLineageRow(
                        lineage_id=self._id("ECL"),
                        export_batch_id=export_batch_id,
                        sheet_name=str(lineage["sheet_name"]),
                        row_index=int(str(lineage["row_index"])),
                        column_index=int(str(lineage["column_index"])),
                        cell_address=str(lineage["cell_address"]),
                        submission_id=str(lineage["submission_id"]),
                        field_key=str(lineage["field_key"]),
                    )
                )
            payload = self._batch(row)
        return payload

    def reexport(
        self,
        *,
        source_batch_id: str,
        idempotency_key: str,
        records: list[dict[str, Any]],
        data_watermark: str,
        actor_id: str,
        reason: str = "",
    ) -> dict[str, object]:
        """Create a re-export that supersedes an existing export batch.

        The source batch is marked SUPERSEDED (but remains downloadable).
        Old file content and hash are immutable.  New lineage rows point to the
        new batch.
        """
        now = self._clock()
        with Session(self._engine) as session, session.begin():
            source = session.get(GovernedExportBatchRow, source_batch_id)
            if source is None:
                raise ReportTemplateError("SOURCE_EXPORT_NOT_FOUND", "源导出批次不存在。")
            if source.status not in ("AVAILABLE", "SUPERSEDED"):
                raise ReportTemplateError(
                    "EXPORT_SOURCE_INVALID", "仅可用或已被替代的导出才能重导。"
                )

            request_payload = {"source_batch_id": source_batch_id, "reason": reason}
            request_hash = self._hash_json(request_payload)

            existing = session.scalar(
                select(GovernedExportBatchRow).where(
                    GovernedExportBatchRow.idempotency_key == idempotency_key
                )
            )
            if existing is not None:
                if existing.request_hash != request_hash:
                    raise ReportTemplateError(
                        "IDEMPOTENCY_CONFLICT", "幂等键对应不同导出请求。"
                    )
                return self._batch(existing)

            template = session.get(ReportTemplateVersionRow, source.template_version_id)
            if template is None:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "模板不存在。")
            workbook = self._workbook(template)
            mapping = session.get(ReportMappingVersionRow, source.mapping_version_id)
            if mapping is None:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "映射不存在。")
            mapping_data = mapping.mapping_json
            sheet = workbook[str(mapping_data["sheet"])]
            start_row = int(mapping_data["start_row"])
            export_batch_id = self._id("GEB")
            lineage_rows: list[dict[str, object]] = []
            for offset, record in enumerate(records):
                row_index = start_row + offset
                submission_id = str(record.get("submission_id", ""))
                for column in mapping_data["columns"]:
                    column_index = int(column["column"])
                    field_key = str(column["source_field"])
                    value = self._safe_cell(record.get(field_key))
                    cell = sheet.cell(row=row_index, column=column_index)
                    cell.value = value
                    lineage_rows.append(
                        {
                            "sheet_name": sheet.title,
                            "row_index": row_index,
                            "column_index": column_index,
                            "cell_address": cell.coordinate,
                            "submission_id": submission_id,
                            "field_key": field_key,
                        }
                    )
            output = BytesIO()
            workbook.save(output)
            file_content = output.getvalue()
            load_workbook(BytesIO(file_content), read_only=True, data_only=False).close()
            new_row = GovernedExportBatchRow(
                export_batch_id=export_batch_id,
                idempotency_key=idempotency_key,
                request_hash=request_hash,
                template_version_id=source.template_version_id,
                mapping_version_id=source.mapping_version_id,
                filters=dict(source.filters or {}),
                data_watermark=data_watermark,
                status="AVAILABLE",
                file_content=file_content,
                file_hash=self._sha(file_content),
                download_name=f"{Path(source.download_name).stem}-RE{now:%Y%m%d%H%M%S}.xlsx",
                created_by=actor_id,
                created_at=now,
                supersedes_batch_id=source_batch_id,
            )
            session.add(new_row)
            session.flush()
            for lineage_entry in lineage_rows:
                session.add(
                    ExportCellLineageRow(
                        lineage_id=self._id("ECL"),
                        export_batch_id=export_batch_id,
                        sheet_name=str(lineage_entry["sheet_name"]),
                        row_index=int(str(lineage_entry["row_index"])),
                        column_index=int(str(lineage_entry["column_index"])),
                        cell_address=str(lineage_entry["cell_address"]),
                        submission_id=str(lineage_entry["submission_id"]),
                        field_key=str(lineage_entry["field_key"]),
                    )
                )
            if source.status == "AVAILABLE":
                source.status = "SUPERSEDED"
                session.add(source)
            payload = self._batch(new_row)
        return payload

    def preview_export(
        self,
        *,
        template_version_id: str,
        mapping_version_id: str,
        filters: dict[str, Any],
        records: list[dict[str, Any]],
    ) -> dict[str, object]:
        """Compute a read-only preview without creating a permanent export record."""
        with Session(self._engine) as session:
            template = session.get(ReportTemplateVersionRow, template_version_id)
            mapping = session.get(ReportMappingVersionRow, mapping_version_id)
            if template is None or mapping is None:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "模板或映射不存在。")
            if mapping.template_version_id != template_version_id:
                raise ReportTemplateError("EXPORT_BINDING_INVALID", "映射不属于所选模板。")
            if mapping.status != "CONFIRMED":
                raise ReportTemplateError(
                    "MAPPING_NOT_CONFIRMED", "未确认映射不能预览。"
                )
        factory_filter = str(filters.get("factory_id", "")).strip() or None
        filtered = records
        if factory_filter:
            filtered = [
                r for r in records
                if str(r.get("factory_id", "")).strip() == factory_filter
            ]
        employees: set[str] = set()
        total_amount = 0.0
        for record in filtered:
            employees.add(str(record.get("subject_employee_code", "")))
            for column in mapping.mapping_json.get("columns", []):
                field_key = str(column.get("source_field", ""))
                value = str(record.get(field_key, "0"))
                try:
                    total_amount += float(value)
                except (ValueError, TypeError):
                    pass
        return {
            "record_count": len(filtered),
            "employee_count": len(employees),
            "total_amount": f"{total_amount:,.2f}",
            "anomaly_count": 0,  # requires anomaly-detection engine; currently not available
        }

    def list_exports(self) -> dict[str, object]:
        with Session(self._engine) as session:
            rows = session.scalars(
                select(GovernedExportBatchRow).order_by(
                    GovernedExportBatchRow.created_at.desc()
                )
            )
            return {"items": [self._batch(row) for row in rows]}

    def download(self, export_batch_id: str) -> bytes:
        with Session(self._engine) as session:
            row = session.get(GovernedExportBatchRow, export_batch_id)
            if row is None:
                raise ReportTemplateError("EXPORT_NOT_AVAILABLE", "导出文件不可下载。")
            if row.status not in ("AVAILABLE", "SUPERSEDED"):
                raise ReportTemplateError("EXPORT_NOT_AVAILABLE", "导出文件不可下载。")
            if self._sha(row.file_content) != row.file_hash:
                raise ReportTemplateError("EXPORT_INTEGRITY_FAILED", "导出文件完整性校验失败。")
            load_workbook(BytesIO(row.file_content), read_only=True).close()
            return bytes(row.file_content)

    def get_export(self, export_batch_id: str) -> dict[str, object]:
        with Session(self._engine) as session:
            row = session.get(GovernedExportBatchRow, export_batch_id)
            if row is None:
                raise ReportTemplateError("EXPORT_NOT_FOUND", "导出批次不存在。")
            return self._batch(row)

    def lineage(self, export_batch_id: str) -> dict[str, object]:
        with Session(self._engine) as session:
            rows = session.scalars(
                select(ExportCellLineageRow)
                .where(ExportCellLineageRow.export_batch_id == export_batch_id)
                .order_by(
                    ExportCellLineageRow.sheet_name,
                    ExportCellLineageRow.row_index,
                    ExportCellLineageRow.column_index,
                )
            )
            return {
                "items": [
                    {
                        "sheet_name": row.sheet_name,
                        "row_index": row.row_index,
                        "column_index": row.column_index,
                        "cell_address": row.cell_address,
                        "submission_id": row.submission_id,
                        "field_key": row.field_key,
                    }
                    for row in rows
                ]
            }

    @staticmethod
    def _check_zip(content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                total = sum(item.file_size for item in archive.infolist())
                compressed = sum(max(item.compress_size, 1) for item in archive.infolist())
                names = {item.filename.lower() for item in archive.infolist()}
        except zipfile.BadZipFile as error:
            raise ReportTemplateError("FILE_DAMAGED", "XLSX 压缩包损坏。") from error
        if total > 50 * 1024 * 1024 or total / max(compressed, 1) > 100:
            raise ReportTemplateError("ZIP_BOMB_RISK", "模板压缩比或解压大小异常。")
        if any("vbaproject" in name or "externallinks/" in name for name in names):
            raise ReportTemplateError("EXTERNAL_CONTENT_FORBIDDEN", "模板含宏或外部链接。")

    @staticmethod
    def _analyze_xlsx(content: bytes) -> tuple[dict[str, Any], list[str]]:
        try:
            workbook = load_workbook(
                BytesIO(content), read_only=False, data_only=False, keep_links=False
            )
        except Exception as error:
            raise ReportTemplateError("FILE_DAMAGED", "无法解析 XLSX 模板。") from error
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            formulas: list[str] = []
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        upper = cell.value.upper()
                        if "DDE(" in upper or "WEBSERVICE(" in upper:
                            raise ReportTemplateError(
                                "DANGEROUS_FORMULA", "模板含危险公式。"
                            )
                        formulas.append(cell.coordinate)
            sheets.append(
                {
                    "name": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                    "formula_cells": formulas,
                    "hidden_rows": [
                        index for index, value in sheet.row_dimensions.items() if value.hidden
                    ],
                    "hidden_columns": [
                        key for key, value in sheet.column_dimensions.items() if value.hidden
                    ],
                    "protected": bool(sheet.protection.sheet),
                }
            )
        workbook.close()
        return {"sheets": sheets}, []

    @staticmethod
    def _analyze_xls(content: bytes) -> tuple[dict[str, Any], list[str]]:
        try:
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
        except Exception as error:
            raise ReportTemplateError("FILE_DAMAGED", "无法解析 XLS 模板。") from error
        sheets = [
            {
                "name": sheet.name,
                "max_row": sheet.nrows,
                "max_column": sheet.ncols,
                "merged_ranges": [
                    f"{get_column_letter(c1 + 1)}{r1 + 1}:"
                    f"{get_column_letter(c2)}{r2}"
                    for r1, r2, c1, c2 in sheet.merged_cells
                ],
                "formula_cells": [],
                "hidden_rows": [],
                "hidden_columns": [],
                "protected": False,
            }
            for sheet in book.sheets()
        ]
        book.release_resources()
        return {"sheets": sheets}, ["旧版 XLS 导出时转换为 XLSX，以确保 WPS/Excel 兼容。"]

    @staticmethod
    def _validate_mapping(
        structure: dict[str, Any], mapping: dict[str, Any]
    ) -> dict[str, Any]:
        names = {str(item["name"]) for item in structure.get("sheets", [])}
        sheet = str(mapping.get("sheet", ""))
        if sheet not in names:
            raise ReportTemplateError("MAPPING_SHEET_INVALID", "映射工作表不存在。")
        start_row = int(mapping.get("start_row", 0))
        columns = mapping.get("columns")
        if start_row < 1 or not isinstance(columns, list) or not columns:
            raise ReportTemplateError("MAPPING_INVALID", "映射起始行或列配置无效。")
        normalized = []
        seen: set[int] = set()
        for item in columns:
            if not isinstance(item, dict):
                raise ReportTemplateError("MAPPING_INVALID", "映射列格式无效。")
            column = int(item.get("column", 0))
            field = str(item.get("source_field", "")).strip()
            if column < 1 or not field or column in seen:
                raise ReportTemplateError("MAPPING_INVALID", "映射列重复或字段为空。")
            seen.add(column)
            normalized.append({"column": column, "source_field": field})
        return {"sheet": sheet, "start_row": start_row, "columns": normalized}

    @staticmethod
    def _workbook(template: ReportTemplateVersionRow):  # type: ignore[no-untyped-def]
        if template.format == "XLSX":
            return load_workbook(
                BytesIO(template.content), read_only=False, data_only=False, keep_links=False
            )
        book = xlrd.open_workbook(file_contents=template.content, formatting_info=False)
        workbook = Workbook()
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)
        for source in book.sheets():
            target = workbook.create_sheet(source.name)
            for row in range(source.nrows):
                for column in range(source.ncols):
                    target.cell(row=row + 1, column=column + 1).value = source.cell_value(
                        row, column
                    )
            for r1, r2, c1, c2 in source.merged_cells:
                target.merge_cells(
                    start_row=r1 + 1,
                    end_row=r2,
                    start_column=c1 + 1,
                    end_column=c2,
                )
        book.release_resources()
        return workbook

    @staticmethod
    def _safe_cell(value: Any) -> Any:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    @staticmethod
    def _sha(content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()

    @staticmethod
    def _hash_json(payload: object) -> str:
        return hashlib.sha256(
            json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()

    @staticmethod
    def _id(prefix: str) -> str:
        return f"{prefix}-{uuid4().hex[:16].upper()}"

    @staticmethod
    def _template(row: ReportTemplateVersionRow) -> dict[str, object]:
        return {
            "template_version_id": row.template_version_id,
            "filename": row.filename,
            "format": row.format,
            "mime_type": row.mime_type,
            "file_hash": row.file_hash,
            "size_bytes": row.size_bytes,
            "structure": row.structure,
            "structure_hash": row.structure_hash,
            "warnings": row.warnings,
            "status": row.status,
            "created_by": row.created_by,
        }

    @staticmethod
    def _mapping(row: ReportMappingVersionRow) -> dict[str, object]:
        return {
            "mapping_version_id": row.mapping_version_id,
            "template_version_id": row.template_version_id,
            "version": row.version,
            "mapping_json": row.mapping_json,
            "content_hash": row.content_hash,
            "status": row.status,
            "created_by": row.created_by,
            "confirmed_by": row.confirmed_by,
        }

    @staticmethod
    def _batch(row: GovernedExportBatchRow) -> dict[str, object]:
        return {
            "export_batch_id": row.export_batch_id,
            "template_version_id": row.template_version_id,
            "mapping_version_id": row.mapping_version_id,
            "filters": row.filters,
            "data_watermark": row.data_watermark,
            "status": row.status,
            "file_hash": row.file_hash,
            "download_name": row.download_name,
            "created_by": row.created_by,
            "supersedes_batch_id": row.supersedes_batch_id or "",
        }
